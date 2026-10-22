"""A sample deck with one slide per layout, for the layout sample page and the eval tooling.

The 17 Oct build plan: "Build the layouts first as a standalone package with a
visual test page that renders every layout with sample content. If a layout
looks wrong with sample content it will look wrong in production."
"""

import io

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from .heuristic import _stat_entries
from .layouts import LAYOUTS
from .plan import _plannable

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def _photo(width=1200, height=800):
    img = Image.new("RGB", (width, height))
    px = img.load()
    for y in range(height):
        t = y / height
        row = (int(20 + 90 * t), int(60 + 120 * t), int(110 + 110 * t))
        for x in range(width):
            px[x, y] = row
    d = ImageDraw.Draw(img)
    d.ellipse((width * 0.62, height * 0.12, width * 0.88, height * 0.52), fill=(250, 235, 200))
    d.polygon([(0, height), (width * 0.35, height * 0.45), (width * 0.7, height)], fill=(12, 30, 50))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def _bullets(tf, items):
    first = True
    for item in items:
        text, level = (item, 0) if isinstance(item, str) else item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.text, p.level, first = text, level, False


def build_sampler_deck(path):
    """A 16:9 deck whose slide n carries content for the n-th layout in ``LAYOUTS`` order."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    L_TITLE, L_CONTENT, L_TWO, L_ONLY, L_BLANK = (prs.slide_layouts[i] for i in (0, 1, 3, 5, 6))

    s = prs.slides.add_slide(L_TITLE)  # 1 title
    s.shapes.title.text = "Safari Fresh Foods"
    s.placeholders[1].text = "Investor update, September 2026"

    s = prs.slides.add_slide(L_TITLE)  # 2 section
    s.shapes.title.text = "Financial performance"
    s.placeholders[1].text = "Revenue, margins and cash"

    s = prs.slides.add_slide(L_CONTENT)  # 3 agenda
    s.shapes.title.text = "Agenda"
    _bullets(s.placeholders[1].text_frame, ["Where we are today", "Market opportunity", "Financial performance",
                                            "Operations and supply chain", "Plans for 2027"])

    s = prs.slides.add_slide(L_CONTENT)  # 4 bullets
    s.shapes.title.text = "The market opportunity"
    _bullets(s.placeholders[1].text_frame, [
        "Urban households spend a growing share of income on fresh produce",
        "Cold-chain coverage outside Nairobi remains thin",
        ("Only 3 of 47 counties have commercial cold storage above 1,000 tonnes", 1),
        ("Farm-gate losses average 30% for leafy greens", 1),
        "Mobile ordering is now mainstream: 71% of customers order via WhatsApp or USSD",
        "Corporate canteens are the fastest-growing channel",
    ])

    s = prs.slides.add_slide(L_TWO)  # 5 two_column
    s.shapes.title.text = "What is working, and what is not"
    _bullets(s.placeholders[1].text_frame, ["Working", ("Repeat orders up 22%", 1),
                                            ("Farmer payments via M-Pesa within 24 hours", 1),
                                            ("Delivery cost per order down to KES 85", 1)])
    _bullets(s.placeholders[2].text_frame, ["Needs fixing", ("Spoilage still 9% in Kisumu hub", 1),
                                            ("Driver churn above 40%", 1),
                                            ("Slow onboarding of corporate clients", 1)])

    s = prs.slides.add_slide(L_CONTENT)  # 6 stats
    s.shapes.title.text = "2026 at a glance"
    _bullets(s.placeholders[1].text_frame, ["KES 450m revenue", "38% gross margin", "12 counties served",
                                            "2,400 farmers onboarded"])

    s = prs.slides.add_slide(L_ONLY)  # 7 table
    s.shapes.title.text = "Financial summary (KES millions)"
    data = [["", "2024", "2025", "2026F"], ["Revenue", "265", "340", "450"], ["Gross profit", "93", "126", "171"],
            ["Operating costs", "(110)", "(128)", "(150)"], ["EBITDA", "(17)", "(2)", "21"],
            ["Cash at year end", "120", "88", "140"]]
    tbl = s.shapes.add_table(6, 4, Inches(0.5), Inches(1.5), Inches(9), Inches(4)).table
    for r, row in enumerate(data):
        for c, v in enumerate(row):
            tbl.cell(r, c).text = v

    s = prs.slides.add_slide(L_ONLY)  # 8 chart
    s.shapes.title.text = "Monthly revenue trend"
    cd = CategoryChartData()
    cd.categories = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep"]
    cd.add_series("2025", (22, 24, 25, 27, 26, 28, 30, 31, 29))
    cd.add_series("2026", (30, 31, 34, 35, 37, 38, 40, 41, 43))
    s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.5), Inches(1.5), Inches(6), Inches(4.5), cd)
    tb = s.shapes.add_textbox(Inches(6.7), Inches(2), Inches(3), Inches(3)).text_frame
    tb.text = "Growth accelerated after the Mombasa hub opened in March."
    tb.add_paragraph().text = "September was our best month ever."

    s = prs.slides.add_slide(L_ONLY)  # 9 image_text
    s.shapes.title.text = "Our new Mombasa hub"
    s.shapes.add_picture(io.BytesIO(_photo()), Inches(0.5), Inches(1.5), Inches(5))
    tb = s.shapes.add_textbox(Inches(5.8), Inches(1.6), Inches(3.8), Inches(4)).text_frame
    _bullets(tb, ["Opened March 2026", "1,200 tonnes cold storage", "Serves 4 coastal counties"])

    s = prs.slides.add_slide(L_ONLY)  # 10 image_full
    s.shapes.title.text = "Harvest season in Nakuru"
    s.shapes.add_picture(io.BytesIO(_photo()), Inches(0.5), Inches(1.4), Inches(9))

    s = prs.slides.add_slide(L_CONTENT)  # 11 timeline
    s.shapes.title.text = "Our journey"
    _bullets(s.placeholders[1].text_frame, ["2023: founded in Nairobi", "2024: first cold-chain hub",
                                            "2025: 12 counties served", "2026: Mombasa hub opened",
                                            "2027: Eldoret and Series A"])

    s = prs.slides.add_slide(L_BLANK)  # 12 quote
    tb = s.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2)).text_frame
    tb.text = "Fresh food should not be a luxury in the cities that grow it."
    tb = s.shapes.add_textbox(Inches(1), Inches(4.5), Inches(8), Inches(1)).text_frame
    tb.text = "Jane Wanjiku, CEO"

    prs.save(path)
    return path


def sampler_plan(deck):
    """The plan that gives slide n of the sampler deck the n-th layout, built from the extracted ids."""
    items = _plannable(deck)
    slides = []
    for layout, s in zip(LAYOUTS, deck.slides):
        own = [items[i] for i in s.order if i in items]
        texts = [x for x in own if x.kind == "text"]
        titles = [t.id for t in texts if t.role == "title"]
        rest = [t for t in texts if t.id not in titles]
        visuals = [x.id for x in own if x.kind != "text"]
        slots = {"title": titles} if titles else {}
        if layout in ("title", "section"):
            slots["subtitle"] = [t.id for t in rest]
        elif layout == "agenda":
            slots["items"] = [t.id for t in rest]
        elif layout == "bullets":
            slots["body"] = [t.id for t in rest]
        elif layout == "two_column":
            boxes = {}
            for t in rest:
                boxes.setdefault(t.box, []).append(t.id)
            left, right = list(boxes.values())[:2]
            slots.update({"left_title": left[:1], "left": left[1:], "right_title": right[:1], "right": right[1:]})
        elif layout == "stats":
            slots["stats"] = [e.to_dict() for e in _stat_entries(rest)]
        elif layout in ("table", "image_full"):
            slots["visuals"] = visuals
        elif layout in ("chart", "image_text"):
            slots["visuals"], slots["body"] = visuals, [t.id for t in rest]
        elif layout == "timeline":
            slots["steps"] = [t.id for t in rest]
        elif layout == "quote":
            ids = [t.id for t in texts]
            slots = {"quote": ids[:1], "attribution": ids[1:]}
        slides.append({"source_slide": s.number, "layout": layout, "slots": {k: v for k, v in slots.items() if v}})
    return {"slides": slides}
