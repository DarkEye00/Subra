"""Step 2: turn a .pptx into structured content (see model.py).

Slides are walked as raw XML, because python-pptx doesn't surface every object:
equations, 3D models and ink sit inside mc:AlternateContent. Anything that must
not be rebuilt (Amendment 1C: video, audio, OLE, equations, 3D models, SmartArt,
ink, and groups holding any of them) becomes a LockedItem with SHA-256 hashes of
all its related parts. Charts are read from the chart XML directly, so unusual
chart types still yield their data.
"""

import copy
import hashlib
import io
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.plot import PlotTypeInspector
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from .model import (
    BackgroundItem,
    ChartItem,
    ChartSeries,
    DeckContent,
    ImageItem,
    LockedItem,
    MasterImage,
    Run,
    SlideContent,
    TableCell,
    TableItem,
    TextItem,
)

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main",
}
A = "{%s}" % NS["a"]
R_NS = NS["r"]
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
OLE_URI = "http://schemas.openxmlformats.org/presentationml/2006/ole"
MODEL3D_URI = "http://schemas.microsoft.com/office/drawing/2017/model3d"
TABLE_URI = "http://schemas.openxmlformats.org/drawingml/2006/table"
CHART_URI = "http://schemas.openxmlformats.org/drawingml/2006/chart"

# Chart types python-pptx can build natively from data.
NATIVE_CHART_TYPES = {
    "AREA", "AREA_STACKED", "AREA_STACKED_100",
    "BAR_CLUSTERED", "BAR_STACKED", "BAR_STACKED_100",
    "COLUMN_CLUSTERED", "COLUMN_STACKED", "COLUMN_STACKED_100",
    "DOUGHNUT", "DOUGHNUT_EXPLODED",
    "LINE", "LINE_MARKERS", "LINE_MARKERS_STACKED", "LINE_MARKERS_STACKED_100",
    "LINE_STACKED", "LINE_STACKED_100",
    "PIE", "PIE_EXPLODED",
    "RADAR", "RADAR_FILLED", "RADAR_MARKERS",
    "XY_SCATTER", "XY_SCATTER_LINES", "XY_SCATTER_LINES_NO_MARKERS",
    "XY_SCATTER_SMOOTH", "XY_SCATTER_SMOOTH_NO_MARKERS",
    "BUBBLE", "BUBBLE_THREE_D_EFFECT",
}

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
BULLETED_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.VERTICAL_OBJECT}
FOOTER_TYPES = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE}
SKIP_TYPES = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.HEADER}

# Relationships we never follow when hashing an object's parts (structure, not content).
STRUCTURAL_RELS = {RT.SLIDE_LAYOUT, RT.SLIDE_MASTER, RT.NOTES_SLIDE, RT.THEME, RT.SLIDE, RT.HYPERLINK,
                   RT.NOTES_MASTER, RT.HANDOUT_MASTER}
# External targets of these relationship types make an object "linked".
LINKED_RELS = {RT.OLE_OBJECT, RT.VIDEO, RT.AUDIO, RT.MEDIA, RT.IMAGE, RT.PACKAGE}

MEDIA_XPATH = ".//a:videoFile | .//a:audioFile | .//a:quickTimeFile | .//p14:media"
MATH_XPATH = ".//a14:m | .//m:oMath | .//m:oMathPara"
LOCKED_XPATH = " | ".join([
    ".//mc:AlternateContent", ".//p:contentPart", MEDIA_XPATH, MATH_XPATH,
    f".//a:graphicData[@uri='{OLE_URI}' or @uri='{DIAGRAM_URI}' or @uri='{MODEL3D_URI}']",
])

DEFAULT_FONT_PT = 18.0


def _xp(el, expr):
    # python-pptx elements override xpath() with a fixed namespace map; use lxml's directly
    return etree._Element.xpath(el, expr, namespaces=NS)


def _local(el):
    return etree.QName(el).localname if isinstance(el.tag, str) else ""


def _sha256(blob):
    return hashlib.sha256(blob).hexdigest()


def extract_deck(path):
    prs = Presentation(str(Path(path)))
    W, H = prs.slide_width, prs.slide_height
    slides = [_extract_slide(n, slide, W, H) for n, slide in enumerate(prs.slides, start=1)]
    return DeckContent(slides=slides, slide_width=W, slide_height=H, unsupported=[],
                       master_images=master_images(prs))


# --- slide walking ----------------------------------------------------------

class _Entry:
    """One content-bearing shape found on the slide, before IDs are assigned."""

    def __init__(self, kind, pos, **kw):
        self.kind = kind  # text | table | chart | image | locked
        self.pos = pos  # (x, y, w, h) in slide EMU
        self.__dict__.update(kw)


def _proxies(shapes, out=None):
    """Map each XML shape element to its python-pptx proxy (recursing into groups)."""
    out = {} if out is None else out
    for shape in shapes:
        out[shape._element] = shape
        if shape._element.tag.endswith("}grpSp"):
            _proxies(shape.shapes, out)
    return out


def _extract_slide(n, slide, W, H):
    entries = []
    _walk(slide.shapes._spTree, _identity, slide.part, _proxies(slide.shapes), entries)
    _detect_title(entries, H)
    entries = _reading_order(entries, H)

    sc = SlideContent(number=n, hidden=slide._element.get("show") in ("0", "false"))
    counters = {"t": 0, "tb": 0, "c": 0, "i": 0, "b": 0, "l": 0, "bg": 0}

    def next_id(kind):
        counters[kind] += 1
        return f"s{n}-{kind}{counters[kind]}"

    for e in entries:
        pos = _frac(e.pos, W, H)
        if e.kind == "text":
            box = next_id("b")
            sc.box_pos[box] = pos
            for para in e.paragraphs:
                item = TextItem(id=next_id("t"), slide=n, role=e.role, box=box, shape_id=e.shape_id, **para)
                sc.texts.append(item)
                sc.order.append(item.id)
            continue
        kind = {"table": "tb", "chart": "c", "image": "i", "locked": "l"}[e.kind]
        item = e.build(next_id(kind), n, pos)
        {"table": sc.tables, "chart": sc.charts, "image": sc.images, "locked": sc.locked}[e.kind].append(item)
        sc.order.append(item.id)

    bg = _background(slide)
    if bg is not None:
        item = BackgroundItem(id=next_id("bg"), slide=n, sha256=bg[0], aspect=bg[1])
        sc.backgrounds.append(item)
        sc.order.append(item.id)

    if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        sc.notes = slide.notes_slide.notes_text_frame.text
    return sc


def _identity(x, y, w, h):
    return x, y, w, h


def _group_transform(grp_el, outer):
    xfrm = grp_el.find(f"{{{NS['p']}}}grpSpPr/{A}xfrm")
    if xfrm is None or xfrm.find(A + "chExt") is None:
        return outer
    off, ext = xfrm.find(A + "off"), xfrm.find(A + "ext")
    ch_off, ch_ext = xfrm.find(A + "chOff"), xfrm.find(A + "chExt")
    ox, oy = int(off.get("x")), int(off.get("y"))
    cx, cy = int(ext.get("cx")), int(ext.get("cy"))
    chx, chy = int(ch_off.get("x")), int(ch_off.get("y"))
    chcx, chcy = int(ch_ext.get("cx")) or 1, int(ch_ext.get("cy")) or 1
    sx, sy = cx / chcx, cy / chcy

    def transform(x, y, w, h):
        return outer(ox + (x - chx) * sx, oy + (y - chy) * sy, w * sx, h * sy)

    return transform


def top_shapes(el):
    """The shape elements that carry an object's position: both branches of mc:AlternateContent."""
    if _local(el) == "AlternateContent":
        return [c for branch in el for c in branch if isinstance(c.tag, str)]
    return [el]


def shape_xfrm(shape_el):
    """The xfrm element holding a shape's own off/ext (None if it inherits from a layout)."""
    tag = _local(shape_el)
    if tag == "grpSp":
        return shape_el.find(f"{{{NS['p']}}}grpSpPr/{A}xfrm")
    if tag == "graphicFrame":
        return shape_el.find(f"{{{NS['p']}}}xfrm")
    if tag == "contentPart":
        return shape_el.find(f"{{{NS['p14']}}}xfrm")
    return shape_el.find(f"{{{NS['p']}}}spPr/{A}xfrm")


def _xfrm_box(xfrm):
    if xfrm is None or xfrm.find(A + "off") is None or xfrm.find(A + "ext") is None:
        return None
    off, ext = xfrm.find(A + "off"), xfrm.find(A + "ext")
    return int(off.get("x", 0)), int(off.get("y", 0)), int(ext.get("cx", 0)), int(ext.get("cy", 0))


def _element_pos(el, transform, proxy=None):
    for shape_el in top_shapes(el):
        box = _xfrm_box(shape_xfrm(shape_el))
        if box:
            return transform(*box)
    if proxy is not None:  # e.g. a placeholder inheriting its position from the layout
        return transform(*[getattr(proxy, a, None) or 0 for a in ("left", "top", "width", "height")])
    return transform(0, 0, 0, 0)


def _walk(container, transform, part, proxies, entries):
    for el in container:
        tag = _local(el)
        if tag in ("", "nvGrpSpPr", "grpSpPr", "extLst"):
            continue
        proxy = proxies.get(el)
        pos = _element_pos(el, transform, proxy)

        if tag in ("AlternateContent", "contentPart"):
            entries.append(_locked_entry(el, classify_locked(el), pos, part))
        elif tag == "grpSp":
            if _xp(el, LOCKED_XPATH):
                entries.append(_locked_entry(el, "group", pos, part))
            else:  # a plain group is dissolved: its text, tables, charts and pictures become items
                _walk(el, _group_transform(el, transform), part, proxies, entries)
        elif tag == "pic":
            if _xp(el, MEDIA_XPATH):
                entries.append(_locked_entry(el, classify_locked(el), pos, part))
            elif not _xp(el, "./p:blipFill/a:blip/@r:embed"):
                entries.append(_locked_entry(el, "linked_image", pos, part))
            else:
                entries.append(_Entry("image", pos, build=_image_builder(proxy, el, part)))
        elif tag == "graphicFrame":
            uri = (_xp(el, "./a:graphic/a:graphicData/@uri") or [""])[0]
            if uri == TABLE_URI and proxy is not None:
                entries.append(_Entry("table", pos, build=_with_shape_id(_table_builder(proxy.table), el)))
            elif uri == CHART_URI and proxy is not None:
                entries.append(_Entry("chart", pos, build=_with_shape_id(_chart_builder(proxy.chart), el)))
            else:
                entries.append(_locked_entry(el, classify_locked(el), pos, part))
        elif tag == "sp":
            if _xp(el, MATH_XPATH):
                entries.append(_locked_entry(el, "equation", pos, part))
                continue
            if proxy is None or not proxy.has_text_frame:
                continue
            role, bulleted, skip = _placeholder_role(proxy)
            if skip:
                continue
            paras = _paragraphs(proxy.text_frame, part, bulleted)
            if paras:
                entries.append(_Entry("text", pos, role=role, paragraphs=paras, shape_id=shape_id_of(el),
                                      font_pt=_max_font_pt(proxy.text_frame), placeholder=proxy.is_placeholder))
        # connectors and decorative shapes without text carry no content


def _with_shape_id(builder, el):
    def build(item_id, slide_no, pos):
        item = builder(item_id, slide_no, pos)
        item.shape_id = shape_id_of(el)
        return item

    return build


# --- locked objects -----------------------------------------------------------

def classify_locked(el):
    if _xp(el, MATH_XPATH):
        return "equation"
    if _xp(el, f".//a:graphicData[@uri='{MODEL3D_URI}']"):
        return "model3d"
    if _local(el) == "contentPart" or _xp(el, ".//p:contentPart"):
        return "ink"
    if _xp(el, f".//a:graphicData[@uri='{OLE_URI}']"):
        return "ole"
    if _xp(el, f".//a:graphicData[@uri='{DIAGRAM_URI}']"):
        return "smartart"
    if _xp(el, ".//a:audioFile"):
        return "audio"
    if _xp(el, MEDIA_XPATH):
        return "video"
    return "other"


def shape_id_of(el):
    ids = etree._Element.xpath(el, ".//*[local-name()='cNvPr']/@id")
    return int(ids[0]) if ids else 0


def related_parts(el, part):
    """(sorted SHA-256 of every part the element depends on, external link targets)."""
    hashes, links, seen = set(), [], set()

    def visit(target):
        if id(target) in seen:
            return
        seen.add(id(target))
        hashes.add(_sha256(target.blob))
        for rel in target.rels.values():
            if not rel.is_external and rel.reltype not in STRUCTURAL_RELS:
                visit(rel.target_part)

    for node in el.iter():
        if not isinstance(node.tag, str):
            continue
        for name, value in node.attrib.items():
            if not name.startswith(f"{{{R_NS}}}") or not value or value not in part.rels:
                continue
            rel = part.rels[value]
            if rel.reltype in STRUCTURAL_RELS:
                continue
            if rel.is_external:
                if rel.reltype in LINKED_RELS and rel.target_ref not in links:
                    links.append(rel.target_ref)
            else:
                visit(rel.target_part)
    return sorted(hashes), links


def xml_signature(el):
    """SHA-256 of the element with positions and relationship IDs blanked (they may legitimately change)."""
    clone = copy.deepcopy(el)
    for node in clone.iter():
        if not isinstance(node.tag, str):
            continue
        if _local(node) in ("off", "ext", "chOff", "chExt"):
            for attr in ("x", "y", "cx", "cy"):
                node.attrib.pop(attr, None)
        for name in list(node.attrib):
            if name.startswith(f"{{{R_NS}}}"):
                node.set(name, "")
    return _sha256(etree.tostring(clone, method="c14n"))


def _locked_entry(el, lock_type, pos, part):
    hashes, links = related_parts(el, part)
    _, _, w, h = pos
    names = etree._Element.xpath(el, ".//*[local-name()='cNvPr']/@name")

    def build(item_id, slide_no, frac_pos):
        return LockedItem(
            id=item_id, slide=slide_no, lock_type=lock_type, shape_id=shape_id_of(el), pos=frac_pos,
            aspect=round(w / h, 4) if w and h else 1.0, part_hashes=hashes, links=links,
            equations=len(_xp(el, ".//m:oMath")), has_media=bool(_xp(el, MEDIA_XPATH)),
            xml_sig=xml_signature(el), name=names[0] if names else "",
        )

    return _Entry("locked", pos, build=build)


# --- backgrounds and master/layout images -----------------------------------------

def _image_size(blob):
    try:
        return Image.open(io.BytesIO(blob)).size
    except Exception:
        return (0, 0)


def _background(slide):
    rid = _xp(slide._element, "./p:cSld/p:bg/p:bgPr/a:blipFill/a:blip/@r:embed")
    if not rid or rid[0] not in slide.part.rels:
        return None
    blob = slide.part.related_part(rid[0]).blob
    w, h = _image_size(blob)
    return _sha256(blob), round(w / h, 4) if w and h else 1.0


def _container_images(el, part, container):
    out = []
    for node in _xp(el, ".//p:pic | ./p:cSld/p:bg"):
        names = etree._Element.xpath(node, ".//*[local-name()='cNvPr']/@name")
        for rid in _xp(node, ".//a:blip/@r:embed | .//asvg:svgBlip/@r:embed"):
            if rid in part.rels and not part.rels[rid].is_external:
                out.append(MasterImage(container=container, sha256=_sha256(part.related_part(rid).blob),
                                       name=names[0] if names else "background"))
    return out


def master_images(prs):
    """Every picture (and picture background) on slide masters and layouts."""
    out = []
    for m, master in enumerate(prs.slide_masters, start=1):
        prefix = "" if m == 1 else f"master{m}-"
        out += _container_images(master._element, master.part, f"master:{m}")
        for i, layout in enumerate(master.slide_layouts, start=1):
            out += _container_images(layout._element, layout.part, f"{prefix}layout:{i}")
    return out


def _placeholder_role(shape):
    """Returns (role, bulleted_by_default, skip)."""
    if not shape.is_placeholder:
        return "body", False, False
    ph_type = shape.placeholder_format.type
    if ph_type in TITLE_TYPES:
        return "title", False, False
    if ph_type == PP_PLACEHOLDER.SUBTITLE:
        return "subtitle", False, False
    if ph_type in FOOTER_TYPES:
        return "footer", False, False
    if ph_type in SKIP_TYPES:
        return "body", False, True
    return "body", ph_type in BULLETED_TYPES, False


# --- text -----------------------------------------------------------------

def _bool_attr(el, name):
    if el is None:
        return None
    v = el.get(name)
    if v is None:
        return None
    return v in ("1", "true")


def _run_from_rpr(text, rpr, part, br=False):
    underline = None
    href = None
    if rpr is not None:
        u = rpr.get("u")
        if u is not None:
            underline = u != "none"
        link = rpr.find(A + "hlinkClick")
        rid = link.get("{%s}id" % NS["r"]) if link is not None else None
        if rid and rid in part.rels and part.rels[rid].is_external:
            href = part.rels[rid].target_ref
    return Run(
        text=text,
        bold=_bool_attr(rpr, "b"),
        italic=_bool_attr(rpr, "i"),
        underline=underline,
        href=href,
        br=br,
    )


def _runs_of(p_el, part):
    runs = []
    for child in p_el:
        tag = etree.QName(child).localname
        if tag in ("r", "fld"):
            t = child.find(A + "t")
            runs.append(_run_from_rpr(t.text or "" if t is not None else "", child.find(A + "rPr"), part))
        elif tag == "br":
            runs.append(_run_from_rpr("\v", child.find(A + "rPr"), part, br=True))
    return runs


def _bullet_of(p_el, bulleted_default):
    ppr = p_el.find(A + "pPr")
    if ppr is not None:
        if ppr.find(A + "buNone") is not None:
            return None, None
        auto = ppr.find(A + "buAutoNum")
        if auto is not None:
            return "number", auto.get("type")
        if ppr.find(A + "buChar") is not None or ppr.find(A + "buBlip") is not None:
            return "char", None
    return ("char" if bulleted_default else None), None


def _paragraphs(text_frame, part, bulleted_default):
    out = []
    for p in text_frame.paragraphs:
        p_el = p._p
        runs = _runs_of(p_el, part)
        text = "".join(r.text for r in runs)
        if not text.strip():
            continue
        bullet, number_type = _bullet_of(p_el, bulleted_default)
        out.append(dict(text=text, runs=runs, level=p.level, bullet=bullet, number_type=number_type))
    return out


def _max_font_pt(text_frame):
    sizes = [int(s) / 100 for s in text_frame._txBody.xpath(".//a:rPr/@sz")]
    return max(sizes) if sizes else None


# --- title detection and reading order ------------------------------------

def _detect_title(entries, H):
    texts = [e for e in entries if e.kind == "text"]
    if any(e.role == "title" for e in texts):
        return
    cands = [
        e for e in texts
        if e.role == "body" and e.pos[1] < 0.35 * H and len(e.paragraphs) <= 2
        and sum(len(p["text"]) for p in e.paragraphs) <= 150
    ]
    if not cands:
        return

    def size(e):
        return e.font_pt or DEFAULT_FONT_PT

    best = max(cands, key=lambda e: (size(e), -e.pos[1]))
    others = [e for e in texts if e is not best]
    bigger = size(best) > max((size(o) for o in others), default=0)
    topmost = all(best.pos[1] < o.pos[1] for o in others)
    if bigger or topmost:
        best.role = "title"


def _reading_order(entries, H):
    titles = [e for e in entries if getattr(e, "role", None) == "title"]
    rest = sorted((e for e in entries if e not in titles), key=lambda e: (e.pos[1], e.pos[0]))
    tolerance = 0.05 * H
    rows, row = [], []
    for e in rest:
        if row and e.pos[1] - row[0].pos[1] > tolerance:
            rows.append(row)
            row = []
        row.append(e)
    if row:
        rows.append(row)
    ordered = titles[:]
    for r in rows:
        ordered.extend(sorted(r, key=lambda e: e.pos[0]))
    return ordered


def _frac(pos, W, H):
    x, y, w, h = pos
    return (round(x / W, 3), round(y / H, 3), round(w / W, 3), round(h / H, 3))


# --- tables ---------------------------------------------------------------

def _table_builder(table):
    rows, cols = len(table.rows), len(table.columns)
    widths = [c.width for c in table.columns]
    total = sum(widths) or 1
    grid = []
    for r in range(rows):
        row = []
        for c in range(cols):
            cell = table.cell(r, c)
            runs = [run for p in cell.text_frame.paragraphs for run in p.runs if run.text.strip()]
            row.append(
                TableCell(
                    row=r,
                    col=c,
                    text=cell.text,
                    bold=bool(runs) and all(run.font.bold for run in runs),
                    row_span=cell.span_height if cell.is_merge_origin else 1,
                    col_span=cell.span_width if cell.is_merge_origin else 1,
                    spanned=cell.is_spanned,
                )
            )
        grid.append(row)

    def build(item_id, slide_no, pos):
        return TableItem(
            id=item_id, slide=slide_no, rows=rows, cols=cols, cells=grid,
            col_widths=[round(w / total, 4) for w in widths],
            header_row=bool(table.first_row), pos=pos,
        )

    return build


# --- charts ---------------------------------------------------------------

def _cache_values(el, numeric):
    """Read a c:cat / c:val / c:xVal / c:yVal / c:bubbleSize element into a list."""
    if el is None:
        return None
    caches = _xp(el, "./c:numRef/c:numCache | ./c:numLit | ./c:strRef/c:strCache | ./c:strLit")
    if not caches:
        multi = _xp(el, "./c:multiLvlStrRef/c:multiLvlStrCache/c:lvl[1]")
        if not multi:
            return None
        cache = multi[0]
        count_el = el.find(".//{%s}ptCount" % NS["c"])
    else:
        cache = caches[0]
        count_el = cache.find("{%s}ptCount" % NS["c"])
    pts = cache.findall("{%s}pt" % NS["c"])
    count = int(count_el.get("val")) if count_el is not None else len(pts)
    out = [None] * count
    for pt in pts:
        idx = int(pt.get("idx"))
        v = pt.find("{%s}v" % NS["c"])
        text = v.text if v is not None else None
        if idx >= count:
            out.extend([None] * (idx + 1 - count))
            count = idx + 1
        if text is None:
            continue
        if numeric:
            try:
                out[idx] = float(text)
            except ValueError:
                out[idx] = None
        else:
            out[idx] = text
    return out


def _format_code(el):
    if el is None:
        return None
    codes = _xp(el, "./c:numRef/c:numCache/c:formatCode/text()")
    return codes[0] if codes else None


def _rich_or_ref_text(title_el):
    if title_el is None:
        return None
    paras = _xp(title_el, "./c:tx/c:rich/a:p")
    if paras:
        text = "\n".join("".join(_xp(p, ".//a:t/text()")) for p in paras)
        return text or None
    ref = _xp(title_el, "./c:tx/c:strRef/c:strCache/c:pt/c:v/text()")
    return ref[0] if ref else None


def _series_name(ser):
    names = _xp(ser, "./c:tx/c:strRef/c:strCache/c:pt/c:v/text() | ./c:tx/c:v/text()")
    return names[0] if names else ""


def _chart_builder(chart):
    space = chart._chartSpace
    plot_area = space.find(".//{%s}plotArea" % NS["c"])
    plot_els = [el for el in plot_area if etree.QName(el).localname.endswith("Chart")]

    try:
        plot_types = [PlotTypeInspector.chart_type(p).name for p in chart.plots]
    except Exception:
        plot_types = []
    chart_type = plot_types[0] if plot_types else "UNKNOWN"

    categories = []
    multi_level = False
    series = []
    for plot_el in plot_els:
        for ser in plot_el.findall("{%s}ser" % NS["c"]):
            cat_el = ser.find("{%s}cat" % NS["c"])
            x_el = ser.find("{%s}xVal" % NS["c"])
            val_el = ser.find("{%s}val" % NS["c"])
            if val_el is None:
                val_el = ser.find("{%s}yVal" % NS["c"])
            if cat_el is not None and not categories:
                multi_level = cat_el.find("{%s}multiLvlStrRef" % NS["c"]) is not None
                categories = [c if c is not None else "" for c in (_cache_values(cat_el, numeric=False) or [])]
            x_values = None
            if x_el is not None:
                x_values = _cache_values(x_el, numeric=True)
                if x_values is not None and all(v is None for v in x_values):
                    # string x values (text labels on a scatter chart): keep as categories
                    categories = [c or "" for c in (_cache_values(x_el, numeric=False) or [])]
                    x_values = None
            series.append(
                ChartSeries(
                    name=_series_name(ser),
                    values=_cache_values(val_el, numeric=True) or [],
                    x_values=x_values,
                    sizes=_cache_values(ser.find("{%s}bubbleSize" % NS["c"]), numeric=True),
                    number_format=_format_code(val_el),
                )
            )

    title = _rich_or_ref_text(space.find("./c:chart/c:title", namespaces=NS))
    axis_titles = {}
    for ax in plot_area:
        if etree.QName(ax).localname not in ("catAx", "valAx", "dateAx", "serAx"):
            continue
        pos_el = ax.find("{%s}axPos" % NS["c"])
        side = "x" if pos_el is not None and pos_el.get("val") in ("b", "t") else "y"
        text = _rich_or_ref_text(ax.find("{%s}title" % NS["c"]))
        if text and side not in axis_titles:
            axis_titles[side] = text

    native = chart_type in NATIVE_CHART_TYPES and len(set(plot_types)) == 1 and not multi_level
    chart_part = chart.part
    workbook = chart_part.chart_workbook.xlsx_part
    chart_xml = etree.tostring(space)
    workbook_blob = workbook.blob if workbook is not None else None
    has_legend = space.find("./c:chart/c:legend", namespaces=NS) is not None

    def build(item_id, slide_no, pos):
        return ChartItem(
            id=item_id, slide=slide_no, chart_type=chart_type, categories=categories,
            series=series, title=title,
            x_axis_title=axis_titles.get("x"), y_axis_title=axis_titles.get("y"),
            has_legend=has_legend, rebuild="native" if native else "keep", pos=pos,
            chart_xml=chart_xml, workbook_blob=workbook_blob,
        )

    return build


# --- images ---------------------------------------------------------------

def _image_builder(shape, el, part):
    image = shape.image
    blob = image.blob
    px_w, px_h = _image_size(blob)
    crop = (shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom)
    src = el.find(f"{{{NS['p']}}}blipFill/{A}srcRect")
    crop_raw = {k: int(src.get(k, 0)) if src is not None else 0 for k in ("l", "t", "r", "b")}
    aspect = (shape.width / shape.height) if shape.width and shape.height else (px_w / px_h if px_h else 1.0)
    cnv = etree._Element.xpath(el, "./*[local-name()='nvPicPr']/*[local-name()='cNvPr']")[0]
    svg = _xp(el, "./p:blipFill/a:blip//asvg:svgBlip/@r:embed")
    svg_sha = _sha256(part.related_part(svg[0]).blob) if svg and svg[0] in part.rels else ""

    def build(item_id, slide_no, pos):
        return ImageItem(
            id=item_id, slide=slide_no, sha256=_sha256(blob), content_type=image.content_type, ext=image.ext,
            px_width=px_w, px_height=px_h, crop=tuple(round(c, 4) for c in crop), crop_raw=crop_raw,
            display_aspect=round(aspect, 4), alt_text=cnv.get("descr", ""), title=cnv.get("title", ""),
            svg_sha256=svg_sha, pos=pos, shape_id=shape_id_of(el), blob=blob,
        )

    return build
