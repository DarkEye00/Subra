"""Stage 6a: deterministic layout checks on the rendered file (17 Oct build plan).

Runs after verification and before any model sees a slide. Cheap rules catch
most faults: text that didn't fit, shapes that overlap or run off the slide,
text below 12pt, poor contrast, and placeholders left empty. Each fault says
whether re-planning could fix it, so the pipeline sends only those slides back
to the planner.
"""

from dataclasses import dataclass, field

from pptx import Presentation
from pptx.util import Inches

from .extract import NS, _local, _xp
from .render import inplace
from .verify import normalize_text

MIN_FONT_PT = 12.0
MIN_CONTRAST = 4.5
OVERLAP_TOLERANCE = 0.02  # share of the smaller shape's area that may be covered
EDGE_TOLERANCE = Inches(0.02)
BACKGROUND_NAMES = ("SUBRA decor", "SUBRA card")
OVERLAY_NAME = "SUBRA overlay"
FOOTER_TYPES = {"ftr", "dt", "sldNum"}
LABEL_CHARS = 30


@dataclass
class Fault:
    kind: str  # overflow | overlap | off_slide | tiny_text | low_contrast | empty_placeholder | vision
    detail: str  # plain words, for the report and for the planner
    repairable: bool = True

    def to_dict(self):
        return {"kind": self.kind, "detail": self.detail, "repairable": self.repairable}


@dataclass
class SlideQA:
    index: int
    source_slide: int
    layout: str
    faults: list = field(default_factory=list)
    repaired: bool = False
    fallback: bool = False
    vision_score: float | None = None

    @property
    def ok(self):
        return not self.faults

    def to_dict(self):
        return {
            "index": self.index, "source_slide": self.source_slide, "layout": self.layout,
            "faults": [f.to_dict() for f in self.faults], "repaired": self.repaired, "fallback": self.fallback,
            "vision_score": self.vision_score,
        }


@dataclass
class QAReport:
    slides: list

    @property
    def faulted(self):
        return [s for s in self.slides if s.faults]

    @property
    def flagged_slides(self):
        return [s.index for s in self.faulted]

    @property
    def repairable_sources(self):
        return sorted({s.source_slide for s in self.slides if any(f.repairable for f in s.faults)})

    def faults_by_source(self):
        """What to tell the planner, per source slide: only the faults re-planning can fix."""
        out = {}
        for s in self.slides:
            for f in s.faults:
                if f.repairable:
                    out.setdefault(s.source_slide, []).append(
                        f"Planned slide {s.index} (layout {s.layout}): {f.detail}"
                    )
        return out

    def mark(self, repaired_sources=(), fallback_sources=()):
        for s in self.slides:
            if s.source_slide in repaired_sources:
                s.repaired = True
            if s.source_slide in fallback_sources:
                s.fallback = True

    def to_dict(self):
        return {"slides": [s.to_dict() for s in self.slides], "flagged_slides": self.flagged_slides}


# --- colour -----------------------------------------------------------------

def _luminance(hex_color):
    channels = []
    for i in (0, 2, 4):
        v = int(hex_color[i:i + 2], 16) / 255
        channels.append(v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4)
    return 0.2126 * channels[0] + 0.7152 * channels[1] + 0.0722 * channels[2]


def contrast_ratio(hex_a, hex_b):
    """WCAG 2.x contrast ratio between two sRGB hex colours (1 to 21)."""
    la, lb = _luminance(hex_a), _luminance(hex_b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


# --- shapes on a rendered slide ------------------------------------------------

@dataclass
class _Shape:
    el: object
    name: str
    kind: str  # text | table | chart | picture | object | group
    box: tuple  # x, y, w, h in EMU
    text: str = ""
    placeholder_type: str | None = None

    @property
    def role(self):
        if self.name.startswith(BACKGROUND_NAMES):
            return "background"
        if self.name.startswith(OVERLAY_NAME):
            return "overlay"
        return "content"

    @property
    def label(self):
        if self.text:
            short = self.text if len(self.text) <= LABEL_CHARS else self.text[:LABEL_CHARS - 1] + "…"
            return f"'{short}' {self.kind}"
        return self.kind

    @property
    def area(self):
        return self.box[2] * self.box[3]


def _text_of(el):
    return normalize_text(" ".join(_xp(el, ".//a:t/text()")))


def _shapes(slide):
    out = []
    for el in slide.shapes._spTree:
        if not isinstance(el.tag, str) or _local(el) in inplace.STRUCTURAL_CHILDREN:
            continue
        box = inplace.element_box(el)
        if box is None:
            continue
        tag = _local(el)
        name = next(iter(_xp(el, ".//*[local-name()='cNvPr']/@name")), "")
        text = _text_of(el)
        placeholder = _xp(el, "./p:nvSpPr/p:nvPr/p:ph")
        ph_type = (placeholder[0].get("type", "body")) if placeholder else None
        if tag == "sp":
            if not text and not name.startswith((*BACKGROUND_NAMES, OVERLAY_NAME)):
                continue  # a shape with no words is decoration (or an empty placeholder, checked separately)
            kind = "text" if text else "shape"
        elif tag == "pic":
            kind = "picture"
        elif tag == "graphicFrame":
            uri = next(iter(_xp(el, ".//a:graphicData/@uri")), "")
            kind = "table" if uri.endswith("/table") else "chart" if uri.endswith("/chart") else "object"
        elif tag == "grpSp":
            kind = "group"
        else:
            kind = "object"
        out.append(_Shape(el, name, kind, box, text, ph_type))
    return out


def _intersection(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    w = min(ax + aw, bx + bw) - max(ax, bx)
    h = min(ay + ah, by + bh) - max(ay, by)
    return w * h if w > 0 and h > 0 else 0


def _contains(box, x, y):
    bx, by, bw, bh = box
    return bx <= x <= bx + bw and by <= y <= by + bh


# --- checks ---------------------------------------------------------------------

def _overflow(entry):
    if not entry.get("overflow"):
        return []
    if entry.get("min_pt", MIN_FONT_PT) < MIN_FONT_PT:
        return [Fault("overflow", "table text is below 12pt")]
    return [Fault("overflow", "text does not fit its box even at 12pt")]


def _overlaps(shapes):
    faults = []
    # decoration, cards and the image_full title band are backgrounds for content, not content
    content = [s for s in shapes if s.role == "content"]
    for i, a in enumerate(content):
        for b in content[i + 1:]:
            shared = _intersection(a.box, b.box)
            smaller = min(a.area, b.area) or 1
            if shared / smaller > OVERLAP_TOLERANCE:
                faults.append(Fault("overlap", f"{a.label} overlaps {b.label}"))
    return faults


def _off_slide(shapes, width, height):
    faults = []
    for s in shapes:
        if s.role == "background":
            continue
        x, y, w, h = s.box
        if x < -EDGE_TOLERANCE or y < -EDGE_TOLERANCE or x + w > width + EDGE_TOLERANCE or y + h > height + EDGE_TOLERANCE:
            faults.append(Fault("off_slide", f"{s.label} runs off the slide"))
    return faults


def _tiny_text(shapes):
    faults = []
    for s in shapes:
        smallest = None
        for run in _xp(s.el, ".//a:r"):
            if not "".join(_xp(run, "./a:t/text()")).strip():
                continue
            sz = next(iter(_xp(run, "./a:rPr/@sz")), None)
            if sz is not None and int(sz) < MIN_FONT_PT * 100:
                smallest = int(sz) if smallest is None else min(smallest, int(sz))
        if smallest is not None:
            faults.append(Fault("tiny_text", f"{smallest / 100:g}pt text in {s.label}", repairable=s.kind != "table"))
    return faults


def _fill_of(el):
    return next(iter(_xp(el, "./p:spPr/a:solidFill/a:srgbClr/@val")), None)


def _slide_background(slide, theme):
    bg = _xp(slide._element, "./p:cSld/p:bg/p:bgPr/a:solidFill/a:srgbClr/@val")
    return bg[0] if bg else theme.bg


def _contrast(slide, shapes, theme):
    faults = []
    backgrounds = [s for s in shapes if s.role in ("background", "overlay") and _fill_of(s.el)]
    slide_bg = _slide_background(slide, theme)
    for s in shapes:
        if s.kind != "text":
            continue
        cx, cy = s.box[0] + s.box[2] / 2, s.box[1] + s.box[3] / 2
        under = [b for b in backgrounds if _contains(b.box, cx, cy)]
        bg = _fill_of(under[-1].el) if under else slide_bg  # the last one drawn is on top
        worst = None
        for run in _xp(s.el, ".//a:r"):
            if not "".join(_xp(run, "./a:t/text()")).strip():
                continue
            color = next(iter(_xp(run, "./a:rPr/a:solidFill/a:srgbClr/@val")), None)
            if color is None:
                continue
            ratio = contrast_ratio(color, bg)
            if ratio < MIN_CONTRAST and (worst is None or ratio < worst[0]):
                worst = (ratio, color)
        if worst:
            faults.append(Fault("low_contrast", f"{s.label} ({worst[1]} on {bg}) has contrast {worst[0]:.1f}:1",
                                repairable=False))
    return faults


def _empty_placeholders(slide):
    faults = []
    for el in slide.shapes._spTree:
        if not isinstance(el.tag, str) or _local(el) != "sp":
            continue
        ph = _xp(el, "./p:nvSpPr/p:nvPr/p:ph")
        if not ph or ph[0].get("type", "body") in FOOTER_TYPES:
            continue
        if not _text_of(el):
            faults.append(Fault("empty_placeholder", f"an empty {ph[0].get('type', 'body')} placeholder"))
    return faults


def check_deck(pptx_path, render_report, theme):
    """One SlideQA per rendered slide, in the order of ``render_report.slides``."""
    prs = Presentation(str(pptx_path))
    width, height = prs.slide_width, prs.slide_height
    out = []
    for slide, entry in zip(prs.slides, render_report.slides):
        qa = SlideQA(index=entry["index"], source_slide=entry["source_slide"], layout=entry["layout"])
        shapes = _shapes(slide)
        qa.faults += _overflow(entry)
        qa.faults += _overlaps(shapes)
        qa.faults += _off_slide(shapes, width, height)
        qa.faults += _tiny_text(shapes)
        qa.faults += _contrast(slide, shapes, theme)
        qa.faults += _empty_placeholders(slide)
        out.append(qa)
    return QAReport(slides=out)
