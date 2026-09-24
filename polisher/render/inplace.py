"""Low-level XML operations for transforming a deck in place (Amendment 1B).

The renderer works on the customer's own file: parts (media, workbooks, models,
images) never leave the package. These helpers find shapes by ID, move kept
elements between slides of the same package, resize locked objects without
touching anything else, and clean up animations that pointed at removed shapes.
"""

from lxml import etree
from pptx.oxml import parse_xml

from ..extract import NS, _local, shape_xfrm, top_shapes

R = "{%s}" % NS["r"]
P = "{%s}" % NS["p"]
A = "{%s}" % NS["a"]
STRUCTURAL_CHILDREN = {"nvGrpSpPr", "grpSpPr", "extLst"}


def _xp(el, expr):
    return etree._Element.xpath(el, expr, namespaces=NS)


# --- finding --------------------------------------------------------------------

def find_shape(sp_tree, shape_id):
    """The shape element with this cNvPr id (the mc:AlternateContent block if it's wrapped in one)."""
    for cnv in sp_tree.iter():
        if not isinstance(cnv.tag, str) or _local(cnv) != "cNvPr" or cnv.get("id") != str(shape_id):
            continue
        shape = cnv.getparent().getparent()
        node = shape
        while node is not None and node is not sp_tree:
            if _local(node) == "AlternateContent":
                return node
            node = node.getparent()
        return shape
    return None


def top_level(el, sp_tree):
    node = el
    while node.getparent() is not sp_tree:
        node = node.getparent()
    return node


def detach_from_group(el, sp_tree):
    """Lift an element out of a (dissolved) group to the top of the shape tree."""
    if el.getparent() is not sp_tree:
        el.getparent().remove(el)
        sp_tree.append(el)


def is_placeholder(el, types):
    ph = _xp(el, "./*[local-name()='nvSpPr']/p:nvPr/p:ph")
    return bool(ph) and ph[0].get("type", "body") in types


# --- relationships ------------------------------------------------------------------

def rids_in(el):
    return {v for node in el.iter() if isinstance(node.tag, str)
            for k, v in node.attrib.items() if k.startswith(R) and v}


def drop_unreferenced(part, candidates):
    """Drop relationships from ``candidates`` that nothing in the part's XML refers to any more."""
    referenced = rids_in(part._element)
    for rid in candidates:
        if rid not in referenced and rid in part.rels:
            part.rels.pop(rid)


def remove_shapes(part, sp_tree, keep):
    """Remove every top-level shape except those in ``keep``. Returns the relationship IDs they used."""
    dropped = set()
    for child in list(sp_tree):
        if not isinstance(child.tag, str) or _local(child) in STRUCTURAL_CHILDREN or child in keep:
            continue
        dropped |= rids_in(child)
        sp_tree.remove(child)
    return dropped


def move_to_slide(el, src_part, dst_slide):
    """Move an element to another slide in the same package, re-pointing its relationships."""
    dst_part = dst_slide.part
    for node in el.iter():
        if not isinstance(node.tag, str):
            continue
        for name, value in list(node.attrib.items()):
            if not name.startswith(R) or not value or value not in src_part.rels:
                continue
            rel = src_part.rels[value]
            target = rel.target_ref if rel.is_external else rel.target_part
            node.set(name, dst_part.relate_to(target, rel.reltype, is_external=rel.is_external))
    moved_rids = rids_in(el)
    el.getparent().remove(el)
    dst_tree = dst_slide.shapes._spTree
    used = {int(i) for i in _xp(dst_tree, ".//*[local-name()='cNvPr']/@id") if str(i).isdigit()}
    own = _xp(el, ".//*[local-name()='cNvPr']")
    if own and own[0].get("id", "").isdigit() and int(own[0].get("id")) in used:
        new_id = str(max(used) + 1)
        for cnv in own:
            if cnv.get("id") == own[0].get("id"):
                cnv.set("id", new_id)
    dst_tree.append(el)
    return moved_rids


# --- geometry -------------------------------------------------------------------------

def element_box(el):
    for shape in top_shapes(el):
        xfrm = shape_xfrm(shape)
        if xfrm is not None and xfrm.find(A + "off") is not None and xfrm.find(A + "ext") is not None:
            off, ext = xfrm.find(A + "off"), xfrm.find(A + "ext")
            return int(off.get("x", 0)), int(off.get("y", 0)), int(ext.get("cx", 0)), int(ext.get("cy", 0))
    return None


def _set_xfrm(xfrm, x, y, w, h):
    off, ext = xfrm.find(A + "off"), xfrm.find(A + "ext")
    if off is None:
        off = etree.SubElement(xfrm, A + "off")
        xfrm.insert(0, off)
    if ext is None:
        ext = etree.SubElement(xfrm, A + "ext")
    off.set("x", str(int(x)))
    off.set("y", str(int(y)))
    ext.set("cx", str(int(w)))
    ext.set("cy", str(int(h)))


def place(el, x, y, w, h, max_scale=None):
    """Fit an element into a box, centred, keeping its aspect ratio. Only off/ext change."""
    box = element_box(el)
    if box is None:
        return None
    _, _, ow, oh = box
    aspect = ow / oh if ow and oh else 1.0
    if w / h > aspect:
        nh, nw = h, h * aspect
    else:
        nw, nh = w, w / aspect
    if max_scale is not None and ow and nw > ow * max_scale:
        nw, nh = ow * max_scale, oh * max_scale
    nx, ny = x + (w - nw) / 2, y + (h - nh) / 2
    for shape in top_shapes(el):
        xfrm = shape_xfrm(shape)
        if xfrm is not None:
            _set_xfrm(xfrm, nx, ny, nw, nh)
    # frames that repeat their position inside (OLE preview pictures, a 3D model's own frame)
    for inner in _xp(el, ".//p:oleObj//p:pic/p:spPr/a:xfrm | .//*[local-name()='model3d']/*[local-name()='spPr']/a:xfrm"):
        _set_xfrm(inner, nx, ny, nw, nh)
    return nx, ny, nw, nh


# --- animations -------------------------------------------------------------------------

def prune_timing(slide_el):
    """Remove animation steps that target shapes no longer on the slide (PowerPoint would flag the file)."""
    timing = slide_el.find(P + "timing")
    if timing is None:
        return
    present = {v for v in _xp(slide_el, "./p:cSld/p:spTree//*[local-name()='cNvPr']/@id")}
    for tgt in _xp(timing, ".//p:spTgt[@spid]"):
        if tgt.get("spid") in present:
            continue
        # remove the whole effect (the p:par whose cTn carries the preset) or media node
        node = tgt
        while node is not None and node is not timing:
            parent = node.getparent()
            if _local(node) == "par" and _xp(node, "./p:cTn[@presetClass or @nodeType]"):
                parent.remove(node)
                break
            if _local(node) in ("video", "audio", "cmd", "set", "anim", "animEffect", "animMotion"):
                parent.remove(node)
                break
            node = parent
    for bld in _xp(timing, "./p:bldLst/*[@spid]"):
        if bld.get("spid") not in present:
            bld.getparent().remove(bld)
    # collapse containers left empty
    changed = True
    while changed:
        changed = False
        for lst in _xp(timing, ".//p:childTnLst[not(*)]"):
            holder = lst.getparent().getparent()  # the p:par / p:seq around this cTn
            if holder is not None and holder.getparent() is not None and _local(holder) in ("par", "seq"):
                holder.getparent().remove(holder)
                changed = True
    bld_lst = timing.find(P + "bldLst")
    if bld_lst is not None and not len(bld_lst):
        timing.remove(bld_lst)
    if not _xp(timing, ".//p:spTgt"):
        slide_el.remove(timing)


# --- slides -----------------------------------------------------------------------------

def add_continuation(prs, after_slide):
    """A new, empty slide with the same layout (placed properly by reorder_slides later)."""
    slide = prs.slides.add_slide(after_slide.slide_layout)
    remove_shapes(slide.part, slide.shapes._spTree, keep=set())
    if after_slide._element.get("show") in ("0", "false"):
        slide._element.set("show", "0")
    return slide


def reorder_slides(prs, ordered):
    sld_id_lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(s.rId): s for s in list(sld_id_lst)}
    new = [by_part[slide.part] for slide in ordered]
    for el in list(sld_id_lst):
        sld_id_lst.remove(el)
    for el in new:
        sld_id_lst.append(el)


def title_placeholder(slide):
    for shape in slide.placeholders:
        if shape.placeholder_format.type in (1, 3):  # TITLE, CENTER_TITLE
            return shape
    return None


def ensure_title(slide):
    """The slide's title placeholder, cloned from its layout if needed (keeps outline view working)."""
    shape = title_placeholder(slide)
    if shape is not None:
        return shape
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.type in (1, 3):
            slide.shapes.clone_placeholder(ph)
            return title_placeholder(slide)
    return None


def reserved_zones(slide):
    """Boxes (EMU) of pictures that show through from the slide's layout and master: logos to avoid."""
    layout = slide.slide_layout
    trees = [layout.shapes._spTree]
    if layout._element.get("showMasterSp") not in ("0", "false"):
        trees.append(layout.slide_master.shapes._spTree)
    zones = []
    for tree in trees:
        for pic in _xp(tree, ".//p:pic"):
            box = element_box(pic)
            if box:
                zones.append(box)
    return zones


def solid_background(slide, hex_color):
    fill = slide.background.fill
    fill.solid()
    from pptx.dml.color import RGBColor

    fill.fore_color.rgb = RGBColor.from_string(hex_color)


def background_picture(slide, rid, name):
    """A picture shape showing the image a slide background used (same image part, no copy)."""
    tree = slide.shapes._spTree
    used = [int(i) for i in _xp(tree, ".//*[local-name()='cNvPr']/@id") if str(i).isdigit()]
    pic = parse_xml(
        f'<p:pic xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}" xmlns:r="{NS["r"]}"><p:nvPicPr>'
        f'<p:cNvPr id="{max(used + [1]) + 1}" name="{name}"/><p:cNvPicPr><a:picLocks noChangeAspect="1"/>'
        f'</p:cNvPicPr><p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/>'
        f'</a:stretch></p:blipFill><p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="100" cy="100"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
    )
    tree.append(pic)
    return pic
