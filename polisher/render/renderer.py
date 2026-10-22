"""Step 4: polish the customer's deck in place (Amendment 1B).

The renderer opens a copy of the uploaded file and transforms it: the theme's
colours and fonts are written into the file's own theme, slide backgrounds are
set, the original text, table, and chart shapes are removed, and redesigned ones
are added per the layout plan. Pictures, locked objects (video, OLE, equations,
SmartArt, ...), and charts python-pptx can't rebuild are kept as the original
elements: only their position and size change, aspect ratio preserved. Masters,
layouts, notes, and animations of kept shapes stay in place.

Layouts are designed on a virtual canvas 7.5 units tall and as wide as the
deck's aspect ratio (13.33 for 16:9, 10 for 4:3), scaled to the real slide.
Text sizes are fitted with real font metrics (textfit.py); slides whose text
still overflows at the minimum size are flagged in the report.
"""

import math
import re
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from ..extract import _xp
from ..verify import normalize_text
from . import inplace
from .textfit import BULLET_INDENT_EM, LEVEL_INDENT_PT, FitPara, fit_size, line_height_ratio, measure, text_width
from .themes import apply_theme, mix

CANVAS_H = 7.5
MX = 0.75  # side margin
TITLE_Y, TITLE_H = 0.4, 1.15
RULE = (MX, 1.68, 0.9, 0.06)
CONTENT_TOP = 2.0
CONTENT_TOP_NO_TITLE = 0.75
CONTENT_BOTTOM = 6.8
GAP = 0.35
PAD = 0.3
MIN_PT = 12.0  # the smallest text size anywhere (17 Oct build plan)
TABLE_FLOOR_PT = 10.0  # a table that can't fit at 12pt may shrink to this and is flagged
FOOTER_PLACEHOLDERS = {"ftr", "dt", "sldNum"}
# locked objects whose content has a fixed text size: never enlarge them
NO_UPSCALE = {"equation", "group", "smartart", "audio"}

NUMERIC_CELL = re.compile(r"^[\s(]*[-+]?(?:KES|KSh|Ksh|USD|\$|€|£)?\s?[\d][\d,.]*\s?(?:%|m|bn|k|x)?\)?\s*$")


class RenderError(RuntimeError):
    pass


@dataclass
class Box:
    x: float
    y: float
    w: float
    h: float

    def inset(self, d):
        return Box(self.x + d, self.y + d, self.w - 2 * d, self.h - 2 * d)

    def intersects(self, o):
        return self.x < o.x + o.w and o.x < self.x + self.w and self.y < o.y + o.h and o.y < self.y + self.h


@dataclass
class TStyle:
    font: str
    color: str
    max_pt: float
    min_pt: float
    bold: bool = False
    italic: bool = False
    override_runs: bool = False  # ignore run-level bold/italic (titles)
    align: object = PP_ALIGN.LEFT
    anchor: object = MSO_ANCHOR.TOP
    line_spacing: float = 1.0
    para_space_em: float = 0.4
    bullets: bool = True
    bullet_color: str | None = None


@dataclass
class RenderReport:
    slides: list = field(default_factory=list)  # [{index, source_slide, layout, overflow, min_pt}]
    dropped_backgrounds: list = field(default_factory=list)  # picture backgrounds the theme replaced
    changes: list = field(default_factory=list)  # human-readable log of notable changes

    @property
    def overflow_slides(self):
        return [s["index"] for s in self.slides if s["overflow"]]

    @property
    def slide_sources(self):
        return [s["source_slide"] for s in self.slides]


def render(plan, deck, theme, output_path, source_path):
    """Transform ``source_path`` (the customer's file, which must be ``deck``'s source) per ``plan``."""
    r = _Renderer(deck, theme, source_path)
    report = r.render(plan)
    r.prs.save(str(output_path))
    return report


def _rgb(hex_color):
    return RGBColor.from_string(hex_color)


class _Renderer:
    def __init__(self, deck, theme, source_path):
        self.deck = deck
        self.t = theme
        self.items = {x.id: x for x in deck.all_items()}
        self.prs = Presentation(str(source_path))
        apply_theme(self.prs, theme)
        self.scale = self.prs.slide_height / Inches(CANVAS_H)
        self.W = CANVAS_H * self.prs.slide_width / self.prs.slide_height  # virtual canvas width
        self.originals = {n: s for n, s in enumerate(self.prs.slides, start=1)}
        self.number_pos = _numbering_positions(deck)
        self.image_side = 0
        self.kept = {}  # item id -> element kept from the original slide
        self.bg_rids = {}  # source slide -> relationship ID of its picture background

    # --- units ----------------------------------------------------------------

    def _in(self, v):
        return int(Inches(v) * self.scale)

    def _emu(self, box):
        return self._in(box.x), self._in(box.y), self._in(box.w), self._in(box.h)

    def _from_emu(self, x, y, w, h):
        f = Inches(1) * self.scale
        return Box(x / f, y / f, w / f, h / f)

    # --- driver ---------------------------------------------------------------

    def render(self, plan):
        report = RenderReport()
        groups = {}
        for ps in plan.slides:
            groups.setdefault(ps.source_slide, []).append(ps)
        missing = set(self.originals) - set(groups)
        if missing:
            raise RenderError(f"plan leaves out source slides {sorted(missing)}")

        ordered, index = [], 0
        placed = set()
        for number in sorted(groups):
            parts = groups[number]
            original = self.originals[number]
            slides, orphans = self._prepare(number, original, parts)
            for ps, slide in zip(parts, slides):
                index += 1
                self.slide, self.source, self.original = slide, number, original
                self.overflow, self.dark, self._back = False, False, 0
                self.min_pt = 999.0
                self.zones = self._zones(slide)
                getattr(self, f"_layout_{ps.layout}")(ps)
                placed.update(ps.ids())
                report.slides.append(
                    {"index": index, "source_slide": number, "layout": ps.layout, "overflow": self.overflow,
                     "min_pt": self.min_pt}
                )
            for slide in slides:
                inplace.prune_timing(slide._element)
            candidates = set(orphans)
            if number in self.bg_rids:  # a replaced picture background no longer references its image
                candidates.add(self.bg_rids[number])
            inplace.drop_unreferenced(original.part, candidates)
            ordered.extend(slides)
        inplace.reorder_slides(self.prs, ordered)

        for s in self.deck.slides:
            for bg in s.backgrounds:
                if bg.id not in placed:
                    report.dropped_backgrounds.append(bg.id)
                    report.changes.append(f"Slide {s.number}: picture background replaced by the theme")
        return report

    def _prepare(self, number, original, parts):
        """Strip the original slide down to what it keeps; add continuation slides for splits."""
        tree = original.shapes._spTree
        for ps in parts:
            for item_id in ps.ids():
                item = self.items[item_id]
                if _keeps_element(item):
                    el = inplace.find_shape(tree, item.shape_id)
                    if el is None:
                        raise RenderError(f"could not find shape {item.shape_id} for {item_id}")
                    if item.kind != "locked":
                        inplace.detach_from_group(el, tree)
                    else:
                        el = inplace.top_level(el, tree)
                    self.kept[item_id] = el
        keep = set(self.kept[i] for ps in parts for i in ps.ids() if i in self.kept)
        if parts[0].slots.get("title"):
            ph = inplace.title_placeholder(original)
            if ph is not None:
                keep.add(ph._element)
        keep |= {el for el in tree if inplace.is_placeholder(el, FOOTER_PLACEHOLDERS)}
        bg = _xp(original._element, "./p:cSld/p:bg/p:bgPr/a:blipFill/a:blip/@r:embed")
        if bg:
            self.bg_rids[number] = bg[0]
        orphans = inplace.remove_shapes(original.part, tree, keep)
        slides = [original] + [inplace.add_continuation(self.prs, original) for _ in parts[1:]]
        return slides, orphans

    def _ids(self, ps, slot):
        return [self.items[i] for i in ps.slots.get(slot, [])]

    # --- slide scaffolding ------------------------------------------------

    def _take_slide(self, bg):
        slide = self.slide
        keep_picture = not self.t.replace_backgrounds and self.source in self.bg_rids and slide is self.original
        if not keep_picture:
            inplace.solid_background(slide, bg)
        return slide

    def _zones(self, slide):
        """Logos and other pictures showing through from the layout or master (small ones only)."""
        zones = []
        for box in inplace.reserved_zones(slide):
            b = self._from_emu(*box)
            if b.w * b.h < 0.15 * self.W * CANVAS_H:
                zones.append(b)
        return zones

    def _avoid(self, box, vertical=False):
        """Shrink a box so it doesn't run into a reserved zone (e.g. a logo on the master)."""
        box = Box(box.x, box.y, box.w, box.h)
        for z in self.zones:
            if not box.intersects(z):
                continue
            if vertical:
                if z.y + z.h / 2 > box.y + box.h / 2:
                    box.h = max(z.y - 0.15 - box.y, box.h * 0.5)
                else:
                    bottom = box.y + box.h
                    box.y = min(z.y + z.h + 0.15, box.y + box.h * 0.5)
                    box.h = bottom - box.y
            elif z.x + z.w / 2 > box.x + box.w / 2:
                box.w = max(z.x - 0.2 - box.x, box.w * 0.5)
            else:
                right = box.x + box.w
                box.x = min(z.x + z.w + 0.2, box.x + box.w * 0.5)
                box.w = right - box.x
        return box

    def _title_shape(self):
        ph = inplace.ensure_title(self.slide)
        return ph

    def _content_slide(self, ps):
        """Standard slide: title at top (in the real title placeholder), content area below."""
        t = self.t
        titles = self._ids(ps, "title")
        slide = self._take_slide(t.bg)
        self._decor_content(slide)
        if titles:
            self._text(
                slide, self._avoid(Box(MX, TITLE_Y, self.W - 2 * MX, TITLE_H)), titles,
                TStyle(t.heading_font, t.title, 32, 20, bold=t.heading_bold, override_runs=True,
                       anchor=MSO_ANCHOR.BOTTOM, line_spacing=0.95, para_space_em=0.1, bullets=False),
                shape=self._title_shape(),
            )
            self._rect(slide, Box(*RULE), t.accent)
            top = CONTENT_TOP
        else:
            top = CONTENT_TOP_NO_TITLE
        return slide, self._avoid(Box(MX, top, self.W - 2 * MX, CONTENT_BOTTOM - top), vertical=True)

    def _body_style(self, max_pt=28, min_pt=MIN_PT, color=None, anchor=MSO_ANCHOR.TOP):
        t = self.t
        return TStyle(t.body_font, color or t.text, max_pt, min_pt, anchor=anchor, bullet_color=t.accent)

    # --- decoration -------------------------------------------------------

    def _rect(self, slide, box, color, shape=MSO_SHAPE.RECTANGLE, radius=None, name="SUBRA decor"):
        if any(box.intersects(z) for z in self.zones):
            return None  # slide shapes draw over the master: never cover a logo with decoration
        sp = slide.shapes.add_shape(shape, *self._emu(box))
        sp.name = name
        sp.fill.solid()
        sp.fill.fore_color.rgb = _rgb(color)
        sp.line.fill.background()
        style = sp._element.find(qn("p:style"))
        if style is not None:
            sp._element.remove(style)
        if radius is not None:
            sp.adjustments[0] = radius
        if name == "SUBRA overlay":  # drawn over the content, e.g. the title band of image_full
            return sp
        # decoration sits behind everything, including the kept pictures and objects
        tree = slide.shapes._spTree
        tree.remove(sp._element)
        tree.insert(list(tree).index(tree.find(qn("p:grpSpPr"))) + 1 + self._back, sp._element)
        self._back += 1
        return sp

    def _card(self, slide, box, strip=True):
        self._rect(slide, box, self.t.surface, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04, name="SUBRA card")
        if strip:
            self._rect(slide, Box(box.x + 0.25, box.y, 0.7, 0.07), self.t.accent)

    def _decor_content(self, slide):
        t = self.t
        if t.decor == "circles":
            self._rect(slide, Box(0, 0, 0.12, CANVAS_H), t.primary)
        elif t.decor == "bands":
            self._rect(slide, Box(0, CANVAS_H - 0.1, self.W, 0.1), t.accent)

    def _decor_cover(self, slide, dark):
        t, W = self.t, self.W
        if t.decor == "circles":
            self._rect(slide, Box(W - 5.03, -1.6, 8.2, 8.2), mix(t.primary, t.on_primary, 0.07), MSO_SHAPE.OVAL)
            self._rect(slide, Box(W - 2.03, 4.9, 1.6, 1.6), t.accent, MSO_SHAPE.OVAL)
        elif t.decor == "bands":
            for i, color in enumerate((t.accent, mix(t.primary, t.accent, 0.5), mix(t.primary, "000000", 0.3))):
                self._rect(slide, Box(0, CANVAS_H - 0.75 + i * 0.25, W, 0.25), color)
        elif t.decor == "block" and not dark:
            self._rect(slide, Box(W * 0.705, 0, W * 0.295, CANVAS_H), t.primary)

    # --- layouts ----------------------------------------------------------

    def _layout_blank(self, ps):
        return self._take_slide(self.t.bg)

    def _layout_title(self, ps):
        return self._cover(ps, section=False)

    def _layout_section(self, ps):
        return self._cover(ps, section=True)

    def _cover(self, ps, section):
        t = self.t
        dark = section or t.cover_on_primary
        self.dark = dark
        slide = self._take_slide(t.primary if dark else t.bg)
        self._decor_cover(slide, dark)
        text_w = min(8.2, self.W * 0.615) if t.decor in ("block", "circles") else self.W - 2 * MX
        titles, subs = self._ids(ps, "title"), self._ids(ps, "subtitle")
        big = (44, 26) if section else (54, 30)
        if titles:
            self._text(
                slide, self._avoid(Box(MX, 1.3, text_w, 3.2)), titles,
                TStyle(t.heading_font, t.on_primary if dark else t.title, *big, bold=t.heading_bold,
                       override_runs=True, anchor=MSO_ANCHOR.BOTTOM, line_spacing=0.95, para_space_em=0.15,
                       bullets=False),
                shape=self._title_shape(),
            )
        self._rect(slide, Box(MX, 4.72, 1.2, 0.08), t.accent)
        if subs:
            self._text(
                slide, self._avoid(Box(MX, 5.0, text_w, 1.5)), subs,
                TStyle(t.body_font, t.on_primary_muted if dark else t.muted, 24, 14, bullets=False,
                       para_space_em=0.25),
            )
        return slide

    def _layout_quote(self, ps):
        t = self.t
        slide = self._take_slide(t.bg)
        self._decor_content(slide)
        self._rect(slide, Box(MX, 1.5, 0.1, 4.0), t.accent)
        width = self.W - 2 * MX - 1.0
        self._text(
            slide, self._avoid(Box(MX + 0.5, 1.3, width, 3.6)), self._ids(ps, "quote"),
            TStyle(t.heading_font, t.title, 48, 20, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0,
                   para_space_em=0.3, bullets=False),
        )
        attribution = self._ids(ps, "attribution")
        if attribution:
            self._text(
                slide, self._avoid(Box(MX + 0.5, 5.1, width, 1.3)), attribution,
                TStyle(t.body_font, t.muted, 20, 12, bullets=False, para_space_em=0.2),
            )
        return slide

    def _layout_bullets(self, ps):
        slide, area = self._content_slide(ps)
        body = self._ids(ps, "body")
        if not body:
            return slide
        style = self._body_style()
        single = self._fit(body, style, area)
        split = _balanced_split(body) if len(body) >= 6 else None
        if split:
            colw = (area.w - GAP * 2) / 2
            boxes = [Box(area.x, area.y, colw, area.h), Box(area.x + colw + GAP * 2, area.y, colw, area.h)]
            parts = [body[:split], body[split:]]
            fits = [self._fit(p, style, b) for p, b in zip(parts, boxes)]
            size = min(f.size for f in fits)
            if size >= single.size + 2 or (not single.fits and all(f.fits for f in fits)):
                for p, b in zip(parts, boxes):
                    self._text(slide, b, p, style, size=size)
                return slide
        self._text(slide, area, body, style)
        return slide

    def _layout_agenda(self, ps):
        t = self.t
        slide, area = self._content_slide(ps)
        items = self._ids(ps, "items")
        if not items:
            return slide
        style = TStyle(t.body_font, t.text, 28, MIN_PT, bullets=False, para_space_em=0.7)
        fit = self._text(slide, Box(area.x, area.y + 0.1, area.w, area.h - 0.1), items, style)
        box = next(sh for sh in slide.shapes if sh.name == "SUBRA text" and sh.text_frame.text.startswith(
            normalize_text(items[0].text)[:20]))
        for p in box.text_frame.paragraphs:
            _agenda_numbering(p, fit.size, t.accent, t.heading_font)
        return slide

    def _layout_timeline(self, ps):
        t = self.t
        slide, area = self._content_slide(ps)
        steps = self._ids(ps, "steps")
        n = len(steps)
        if not n:
            return slide
        colw = area.w / n
        line_y = area.y + 0.9
        self._rect(slide, Box(area.x + colw / 2, line_y, colw * (n - 1), 0.06), t.line)
        for i in range(n):
            cx = area.x + colw / 2 + i * colw
            self._rect(slide, Box(cx - 0.18, line_y - 0.15, 0.36, 0.36), t.accent, MSO_SHAPE.OVAL)
        style = TStyle(t.body_font, t.text, 20, MIN_PT, bullets=False, align=PP_ALIGN.CENTER, para_space_em=0.2)
        boxes = [Box(area.x + i * colw + 0.1, line_y + 0.45, colw - 0.2, area.h - 1.5) for i in range(n)]
        size = min(self._fit([step], style, box).size for step, box in zip(steps, boxes))
        for step, box in zip(steps, boxes):
            self._text(slide, box, [step], style, size=size)
        return slide

    def _layout_image_full(self, ps):
        t = self.t
        slide = self._take_slide(t.bg)
        band_h = 1.3
        self._place_visuals(slide, self._ids(ps, "visuals"), Box(0, 0, self.W, CANVAS_H - band_h), stacked=False)
        band = self._rect(slide, Box(0, CANVAS_H - band_h, self.W, band_h), t.primary, name="SUBRA overlay")
        titles = self._ids(ps, "title")
        if titles:
            color = t.on_primary if band is not None else t.title  # no band when a logo sits there
            shape = self._title_shape()
            self._text(
                slide, self._avoid(Box(MX, CANVAS_H - band_h + 0.2, self.W - 2 * MX, band_h - 0.4)), titles,
                TStyle(t.heading_font, color, 32, 20, bold=t.heading_bold, override_runs=True,
                       anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95, para_space_em=0.1, bullets=False),
                shape=shape,
            )
            if shape is not None:  # the placeholder came before the picture and the band; show it on top
                tree = slide.shapes._spTree
                tree.remove(shape._element)
                tree.append(shape._element)
        return slide

    def _layout_two_column(self, ps):
        t = self.t
        slide, area = self._content_slide(ps)
        colw = (area.w - GAP) / 2
        heading_style = TStyle(t.heading_font, t.highlight, 24, 16, bold=True, bullets=False, para_space_em=0.1,
                               anchor=MSO_ANCHOR.TOP)
        body_style = self._body_style(24, 12)
        columns = []
        for i, side in enumerate(("left", "right")):
            card = Box(area.x + i * (colw + GAP), area.y, colw, area.h)
            self._card(slide, card)
            inner = card.inset(PAD)
            inner = Box(inner.x, inner.y + 0.1, inner.w, inner.h - 0.1)
            heads, body = self._ids(ps, f"{side}_title"), self._ids(ps, side)
            used = 0.0
            if heads:
                hbox = Box(inner.x, inner.y, inner.w, min(1.2, inner.h * 0.3))
                fit = self._text(slide, hbox, heads, heading_style)
                used = fit.height / 72 / self.scale + 0.2
            columns.append((Box(inner.x, inner.y + used, inner.w, inner.h - used), body))
        fitted = [self._fit(body, body_style, box).size for box, body in columns if body]
        size = min(fitted) if fitted else None
        for box, body in columns:
            if body:
                self._text(slide, box, body, body_style, size=size)
        return slide

    def _layout_stats(self, ps):
        t = self.t
        slide, area = self._content_slide(ps)
        entries = ps.slots.get("stats", [])
        body = self._ids(ps, "body")
        parts = [self._stat_parts(e) for e in entries]
        values = ["".join(r.text for r in runs) for runs, _ in parts]
        stats_area = Box(area.x, area.y, area.w, area.h * (0.62 if body else 1.0))

        def value_size_for(cols):
            inner_w = ((stats_area.w - GAP * (cols - 1)) / cols - 2 * PAD) * self.scale
            size = 60.0
            while size > 24 and any(
                text_width(v, t.heading_font, size, bold=True) > inner_w * 72 * 0.94 for v in values
            ):
                size -= 2
            return size

        n = max(len(entries), 1)
        cols, rows = n, 1
        value_size = value_size_for(n)
        if n == 4 and value_size < 40:
            cols, rows = 2, 2
            value_size = value_size_for(2)
        card_w = (stats_area.w - GAP * (cols - 1)) / cols
        max_card_h = (stats_area.h - GAP * (rows - 1)) / rows
        value_h = value_size * line_height_ratio(t.heading_font) / 72 / self.scale + 0.1
        label_style = TStyle(t.body_font, t.text, 20, MIN_PT, bullets=False, para_space_em=0.25)
        label_box = Box(0, 0, card_w - 2 * PAD, max_card_h - 2 * PAD - 0.15 - value_h)
        label_fits = [self._fit(labels, label_style, label_box) for _, labels in parts if labels]
        label_size = min((f.size for f in label_fits), default=16)
        label_h = max((f.height for f in label_fits), default=0) / 72 / self.scale
        card_h = min(max(value_h + label_h + 2 * PAD + 0.35, 1.9), max_card_h)
        block_h = card_h * rows + GAP * (rows - 1)
        top = stats_area.y + (0 if body else (stats_area.h - block_h) * 0.45)
        cards = [
            Box(stats_area.x + (i % cols) * (card_w + GAP), top + (i // cols) * (card_h + GAP), card_w, card_h)
            for i in range(n)
        ]
        for card in cards:
            self._card(slide, card)
        for (value_runs, labels), card, entry in zip(parts, cards, entries):
            inner = card.inset(PAD)
            inner = Box(inner.x, inner.y + 0.15, inner.w, inner.h - 0.15)
            self._write_stat(slide, inner, value_runs, labels, entry, value_size, label_size)
        if body:
            body_top = top + block_h + 0.35
            self._text(slide, Box(area.x, body_top, area.w, area.y + area.h - body_top), body,
                       self._body_style(20, 12))
        return slide

    def _stat_parts(self, entry):
        """Returns (value_runs, label_items). Label items may include a synthetic tail of a highlight split."""
        items = [self.items[i] for i in entry.ids]
        first = items[0]
        if entry.highlight:
            head, tail = _split_runs(first, entry.highlight)
            labels = ([_TailItem(tail, first)] if tail else []) + items[1:]
            return head, labels
        return first.runs, items[1:]

    def _write_stat(self, slide, box, value_runs, labels, entry, value_size, label_size):
        t = self.t
        shape = slide.shapes.add_textbox(*self._emu(box))
        shape.name = "SUBRA text"
        self.min_pt = min(self.min_pt, value_size, label_size)
        tf = shape.text_frame
        _prepare_frame(tf, MSO_ANCHOR.TOP, fits=True)
        tf.clear()
        p = tf.paragraphs[0]
        _no_bullet(p, 0)
        for run in value_runs:
            self._run(p, run, t.heading_font, value_size, t.highlight, bold=True, italic=False, override=True)
        rest = labels
        if entry.highlight and labels and isinstance(labels[0], _TailItem):
            p._p.add_br()
            for run in labels[0].runs:
                self._run(p, run, t.body_font, label_size, t.text, bold=False, italic=False)
            rest = labels[1:]
        for item in rest:
            q = tf.add_paragraph()
            q.space_before = Pt(label_size * 0.4)
            _no_bullet(q, 0)
            for run in item.runs:
                self._run(q, run, t.body_font, label_size, t.text, bold=False, italic=False)

    def _layout_table(self, ps):
        slide, area = self._content_slide(ps)
        visuals, body = self._ids(ps, "visuals"), self._ids(ps, "body")
        if body:
            style = self._body_style(18, MIN_PT)
            body_h = min(self._fit(body, style, Box(0, 0, area.w, 1.4)).height / 72 / self.scale + 0.1, 1.4)
            self._text(slide, Box(area.x, area.y + area.h - body_h, area.w, body_h), body, style)
            area = Box(area.x, area.y, area.w, area.h - body_h - 0.3)
        self._place_visuals(slide, visuals, area, stacked=False)
        return slide

    def _layout_chart(self, ps):
        slide, area = self._content_slide(ps)
        visuals, body = self._ids(ps, "visuals"), self._ids(ps, "body")
        if body:
            vis_w = area.w * 0.63
            panel = Box(area.x + vis_w + GAP, area.y, area.w - vis_w - GAP, area.h)
            self._card(slide, panel)
            self._text(slide, panel.inset(PAD), body, self._body_style(20, MIN_PT, anchor=MSO_ANCHOR.MIDDLE))
            area = Box(area.x, area.y, vis_w, area.h)
        self._place_visuals(slide, visuals, area, stacked=bool(body))
        return slide

    def _layout_image_text(self, ps):
        slide, area = self._content_slide(ps)
        visuals, body = self._ids(ps, "visuals"), self._ids(ps, "body")
        if body:
            half = (area.w - GAP) / 2
            left_images = self.image_side % 2 == 0
            self.image_side += 1
            img_box = Box(area.x if left_images else area.x + half + GAP, area.y, half, area.h)
            txt_box = Box(area.x + half + GAP if left_images else area.x, area.y, half, area.h)
            self._text(slide, txt_box, body, self._body_style(26, 12, anchor=MSO_ANCHOR.MIDDLE))
            area = img_box
        self._place_visuals(slide, visuals, area, stacked=False)
        return slide

    # --- visuals ----------------------------------------------------------

    def _place_visuals(self, slide, visuals, area, stacked):
        n = len(visuals)
        if n == 0:
            return
        if n == 1:
            cells = [area]
        elif n == 2 and stacked:
            h = (area.h - GAP) / 2
            cells = [Box(area.x, area.y, area.w, h), Box(area.x, area.y + h + GAP, area.w, h)]
        elif n == 2:
            w = (area.w - GAP) / 2
            cells = [Box(area.x, area.y, w, area.h), Box(area.x + w + GAP, area.y, w, area.h)]
        else:
            cols = 2
            rows = math.ceil(n / cols)
            w, h = (area.w - GAP) / cols, (area.h - GAP * (rows - 1)) / rows
            cells = [Box(area.x + (i % cols) * (w + GAP), area.y + (i // cols) * (h + GAP), w, h) for i in range(n)]
        for item, cell in zip(visuals, cells):
            if item.kind == "table":
                self._table(slide, item, cell)
            elif item.kind == "chart" and not _keeps_element(item):
                self._chart(slide, item, cell)
            elif item.kind == "background":
                self._background_picture(slide, item, cell)
            else:  # pictures, locked objects, charts we keep as they are
                self._place_kept(slide, item, cell)

    def _place_kept(self, slide, item, cell):
        el = self.kept[item.id]
        if slide is not self.original:  # a split slide: move the picture (never a locked object) across
            orphans = inplace.move_to_slide(el, self.original.part, slide)
            inplace.drop_unreferenced(self.original.part, orphans)
        max_scale = 1.0 if item.kind == "locked" and item.lock_type in NO_UPSCALE else None
        inplace.place(el, *self._emu(cell), max_scale=max_scale)

    def _background_picture(self, slide, item, cell):
        rid = self.bg_rids[item.slide]
        if slide is not self.original:
            rid = slide.part.relate_to(self.original.part.related_part(rid), RT.IMAGE)
        pic = inplace.background_picture(slide, rid, "Background picture")
        ext = pic.find(f"{inplace.P}spPr/{inplace.A}xfrm/{inplace.A}ext")
        ext.set("cx", str(int(1_000_000 * item.aspect)))
        ext.set("cy", "1000000")
        inplace.place(pic, *self._emu(cell))

    def _table(self, slide, item, cell):
        t = self.t
        widths = [cell.w * f for f in item.col_widths]
        size, fits = TABLE_FLOOR_PT, False
        for candidate in range(18, int(TABLE_FLOOR_PT) - 1, -1):
            heights = self._row_heights(item, widths, candidate)
            if sum(heights) <= cell.h:
                size, fits = float(candidate), True
                break
        heights = self._row_heights(item, widths, size)
        if fits:  # breathe: grow rows toward a comfortable height while the table still fits
            comfy = min(0.55, cell.h / item.rows)
            heights = [max(h, comfy) for h in heights]
        if not fits or size < MIN_PT:  # the table only fits below the 12pt floor, or not at all
            self.overflow = True
        self.min_pt = min(self.min_pt, size)
        gf = slide.shapes.add_table(item.rows, item.cols, self._in(cell.x), self._in(cell.y), self._in(cell.w),
                                    self._in(sum(heights)))
        gf.name = "SUBRA table"
        table = gf.table
        table.first_row = item.header_row
        table.horz_banding = False
        for c, w in enumerate(widths):
            table.columns[c].width = self._in(w)
        for r, h in enumerate(heights):
            table.rows[r].height = self._in(h)
        for row in item.cells:
            for c in row:
                if c.row_span > 1 or c.col_span > 1:
                    table.cell(c.row, c.col).merge(table.cell(c.row + c.row_span - 1, c.col + c.col_span - 1))
        numeric_cols = self._numeric_columns(item)
        for row in item.cells:
            for c in row:
                tc = table.cell(c.row, c.col)
                header = item.header_row and c.row == 0
                if header:
                    fill, color, bold = t.primary, t.on_primary, True
                else:
                    band = (c.row - (1 if item.header_row else 0)) % 2 == 0
                    fill, color, bold = (t.surface if band else t.bg), t.text, c.bold
                tc.fill.solid()
                tc.fill.fore_color.rgb = _rgb(fill)
                tc.margin_left = tc.margin_right = Inches(0.08)
                tc.margin_top = tc.margin_bottom = Inches(0.04)
                tc.vertical_anchor = MSO_ANCHOR.MIDDLE
                align = PP_ALIGN.RIGHT if numeric_cols[c.col] and (header or NUMERIC_CELL.match(c.text)) else PP_ALIGN.LEFT
                self._plain_text(tc.text_frame, c.text, t.body_font, size, color, bold, align)

    def _numeric_columns(self, item):
        body = item.cells[1:] if item.header_row else item.cells
        out = []
        for c in range(item.cols):
            vals = [row[c].text for row in body if not row[c].spanned and row[c].text.strip()]
            out.append(bool(vals) and sum(bool(NUMERIC_CELL.match(v)) for v in vals) / len(vals) >= 0.6)
        return out

    def _row_heights(self, item, widths, size):
        margin_pt = 0.08 * 2 * 72
        ratio = line_height_ratio(self.t.body_font)
        heights = [0.0] * item.rows
        for row in item.cells:
            for c in row:
                if c.spanned:
                    continue
                w = sum(widths[c.col:c.col + c.col_span]) * 72 * self.scale - margin_pt
                paras = [FitPara(runs=[(line, c.bold or c.row == 0, False)]) for line in (c.text or " ").split("\n")]
                need = measure(paras, self.t.body_font, size, w, para_space_em=0).height + 0.08 * 72
                need = max(need, size * ratio + 0.08 * 72) / c.row_span
                for r in range(c.row, c.row + c.row_span):
                    heights[r] = max(heights[r], need)
        return [h / 72 / self.scale for h in heights]

    def _plain_text(self, tf, text, font, size, color, bold, align):
        tf.clear()
        tf.word_wrap = True
        for i, line in enumerate((text or "").split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            parts = line.split("\v")
            for j, part in enumerate(parts):
                if j:
                    p._p.add_br()
                if part:
                    r = p.add_run()
                    r.text = part
                    f = r.font
                    f.size, f.name, f.bold = Pt(size), font, bool(bold)
                    f.color.rgb = _rgb(color)

    def _chart(self, slide, item, cell):
        gf = slide.shapes.add_chart(XL_CHART_TYPE[item.chart_type], *self._emu(cell), _chart_data(item))
        gf.name = "SUBRA chart"
        self._style_chart(gf.chart, item)

    def _style_chart(self, chart, item):
        t = self.t
        palette = [_rgb(c) for c in t.chart]
        chart.font.size = Pt(14)
        chart.font.name = t.body_font
        chart.font.color.rgb = _rgb(t.text)
        if item.title:
            chart.has_title = True
            tf = chart.chart_title.text_frame
            tf.text = item.title
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.size, r.font.bold, r.font.name = Pt(16), True, t.body_font
                    r.font.color.rgb = _rgb(t.text)
        else:
            chart.has_title = False
        family = item.chart_type.split("_")[0]
        plot = chart.plots[0]
        if family in ("PIE", "DOUGHNUT"):
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            for i, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = palette[i % len(palette)]
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.show_value, labels.show_percentage, labels.show_category_name = False, True, False
            labels.number_format, labels.number_format_is_linked = "0%", False
            labels.position = XL_LABEL_POSITION.OUTSIDE_END
            labels.font.size, labels.font.bold = Pt(14), True
            labels.font.color.rgb = _rgb(t.text)
            return
        chart.has_legend = len(item.series) > 1 or item.has_legend
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        markers_only = item.chart_type == "XY_SCATTER"
        for i, series in enumerate(plot.series):
            color = palette[i % len(palette)]
            if family in ("LINE", "RADAR") or (family == "XY" and not markers_only):
                series.format.line.color.rgb = color
                series.format.line.width = Pt(2.5)
            if (family in ("LINE", "RADAR") and "MARKERS" in item.chart_type) or (
                family == "XY" and "NO_MARKERS" not in item.chart_type
            ):
                series.marker.style = XL_MARKER_STYLE.CIRCLE
                series.marker.size = 9
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = color
                series.marker.format.line.color.rgb = color
            if family in ("BAR", "COLUMN", "AREA", "BUBBLE"):
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = color
        if family in ("BAR", "COLUMN"):
            plot.gap_width = 70
        try:
            value_axis, category_axis = chart.value_axis, chart.category_axis
        except ValueError:
            return
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = _rgb(t.line)
        value_axis.format.line.fill.background()
        category_axis.format.line.color.rgb = _rgb(t.line)
        for axis in (value_axis, category_axis):
            axis.tick_labels.font.size = Pt(13)
            axis.tick_labels.font.color.rgb = _rgb(t.muted)
        for axis, text in ((category_axis, item.x_axis_title), (value_axis, item.y_axis_title)):
            if text:
                axis.has_title = True
                axis.axis_title.text_frame.text = text

    # --- text -------------------------------------------------------------

    def _fit(self, items, style, box):
        base_level = min((item.level for item in items), default=0)
        paras = [
            FitPara(
                runs=[(r.text, style.bold if style.override_runs else bool(r.bold if r.bold is not None else style.bold),
                       bool(r.italic)) for r in item.runs],
                level=item.level - base_level,
                bullet=bool(style.bullets and item.bullet),
            )
            for item in items
        ]
        return fit_size(paras, style.font, box.w * 72 * self.scale, box.h * 72 * self.scale, style.max_pt,
                        style.min_pt, line_spacing=style.line_spacing, para_space_em=style.para_space_em)

    def _text(self, slide, box, items, style, shape=None, size=None):
        fit = self._fit(items, style, box)
        if size is not None:
            fit.size = size
            fit.fits = self._fit(items, TStyle(**{**style.__dict__, "max_pt": size, "min_pt": size}), box).fits
        if not fit.fits:
            self.overflow = True
        self.min_pt = min(self.min_pt, fit.size)
        if shape is None:
            shape = slide.shapes.add_textbox(*self._emu(box))
            shape.name = "SUBRA text"
        else:
            shape.left, shape.top, shape.width, shape.height = self._emu(box)
        tf = shape.text_frame
        _prepare_frame(tf, style.anchor, fit.fits)
        tf.clear()
        last_pos = {}
        base_level = min((item.level for item in items), default=0)
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = style.align
            p.line_spacing = style.line_spacing
            if i:
                p.space_before = Pt(fit.size * style.para_space_em)
            bullet = item.bullet if style.bullets else None
            start_at = None
            if bullet == "number":
                pos = self.number_pos.get(item.id, 1)
                if pos != last_pos.get(item.level, 0) + 1:
                    start_at = pos
                last_pos[item.level] = pos
            _set_bullet(p, item.level - base_level, bullet, item.number_type, fit.size, style.bullet_color, start_at)
            for run in item.runs:
                self._run(p, run, style.font, fit.size, style.color, style.bold, style.italic, style.override_runs)
        return fit

    def _run(self, p, run, font, size, color, bold, italic, override=False):
        if run.br:
            p._p.add_br()
            return
        if not run.text:
            return
        r = p.add_run()
        r.text = run.text
        f = r.font
        f.size = Pt(size)
        f.name = font
        f.color.rgb = _rgb(color)
        if override:
            f.bold, f.italic = bold, italic
        else:
            f.bold = run.bold if run.bold is not None else bold
            f.italic = run.italic if run.italic is not None else italic
        if run.underline:
            f.underline = True
        if run.href:
            r.hyperlink.address = run.href


class _TailItem:
    """The label part of a highlighted stat paragraph (not a separate content item)."""

    def __init__(self, runs, source):
        self.runs = runs
        self.level = 0
        self.bullet = None
        self.text = "".join(r.text for r in runs)
        self.source = source


def _split_runs(item, highlight):
    """Split a paragraph's runs after the highlight prefix; drop the whitespace between."""
    from ..model import Run

    text = item.text
    target = normalize_text(highlight)
    cut = next((k for k in range(1, len(text) + 1) if normalize_text(text[:k]) == target), len(text))
    head, tail, pos = [], [], 0
    for run in item.runs:
        end = pos + len(run.text)
        if end <= cut:
            head.append(run)
        elif pos >= cut:
            tail.append(run)
        else:
            k = cut - pos
            head.append(Run(**{**run.__dict__, "text": run.text[:k]}))
            tail.append(Run(**{**run.__dict__, "text": run.text[k:]}))
        pos = end
    while tail and (tail[0].br or not tail[0].text.strip()):
        tail.pop(0)
    if tail:
        tail[0] = Run(**{**tail[0].__dict__, "text": tail[0].text.lstrip()})
    return head, tail


def _prepare_frame(tf, anchor, fits):
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.auto_size = MSO_AUTO_SIZE.NONE if fits else MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _agenda_numbering(p, size, color, font):
    """Large accent numbers in the heading font; auto-numbers are bullets, so they add no text to the deck."""
    _set_bullet(p, 0, "number", "arabicPeriod", size, color, None)
    pPr = p._p.get_or_add_pPr()
    hang = Pt(size * 2.2)
    pPr.set("marL", str(hang))
    pPr.set("indent", str(-hang))
    clr = pPr.find(qn("a:buClr"))
    anchor = clr if clr is not None else pPr.find(qn("a:buAutoNum"))
    sz = pPr.makeelement(qn("a:buSzPct"), {"val": "150000"})
    bu_font = pPr.makeelement(qn("a:buFont"), {"typeface": font})
    if clr is not None:
        clr.addnext(sz)
    else:
        anchor.addprevious(sz)
    sz.addnext(bu_font)


def _no_bullet(p, level):
    _set_bullet(p, level, None, None, 0, None, None)


def _set_bullet(p, level, bullet, number_type, size, color, start_at):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buNone", "a:buAutoNum", "a:buChar"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if level:
        pPr.set("lvl", str(level))
    base = level * LEVEL_INDENT_PT
    if bullet:
        hang = BULLET_INDENT_EM * size * (1.4 if bullet == "number" else 1.0)
        pPr.set("marL", str(Pt(base + hang)))
        pPr.set("indent", str(-Pt(hang)))
        if color:
            clr = pPr.makeelement(qn("a:buClr"), {})
            clr.append(clr.makeelement(qn("a:srgbClr"), {"val": color}))
            _append_bullet_el(pPr, clr)
        if bullet == "number":
            attrs = {"type": number_type or "arabicPeriod"}
            if start_at and start_at > 1:
                attrs["startAt"] = str(start_at)
            _append_bullet_el(pPr, pPr.makeelement(qn("a:buAutoNum"), attrs))
        else:
            _append_bullet_el(pPr, pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
            _append_bullet_el(pPr, pPr.makeelement(qn("a:buChar"), {"char": "•" if level == 0 else "–"}))
    else:
        pPr.set("marL", str(Pt(base)))
        pPr.set("indent", "0")
        _append_bullet_el(pPr, pPr.makeelement(qn("a:buNone"), {}))


def _append_bullet_el(pPr, el):
    # bullet properties go after spacing and before tabLst/defRPr/extLst
    successor = next((c for c in pPr if c.tag in (qn("a:tabLst"), qn("a:defRPr"), qn("a:extLst"))), None)
    if successor is not None:
        successor.addprevious(el)
    else:
        pPr.append(el)


def _keeps_element(item):
    """Pictures, locked objects, and charts we can't rebuild keep their original element."""
    if item.kind in ("image", "locked"):
        return True
    return item.kind == "chart" and (item.rebuild != "native" or _chart_data(item) is None)


def _chart_data(item):
    """python-pptx chart data for a native rebuild, or None if the chart should be kept as it is."""
    family = item.chart_type.split("_")[0]
    if not item.series:
        return None
    if family == "XY":
        data = XyChartData()
        for s in item.series:
            if s.x_values is None or len(s.x_values) != len(s.values) or None in s.x_values or None in s.values:
                return None
            ser = data.add_series(s.name, number_format=s.number_format)
            for x, y in zip(s.x_values, s.values):
                ser.add_data_point(x, y)
        return data
    if family == "BUBBLE":
        data = BubbleChartData()
        for s in item.series:
            if not s.x_values or not s.sizes or None in s.x_values or None in s.values or None in s.sizes:
                return None
            ser = data.add_series(s.name, number_format=s.number_format)
            for x, y, z in zip(s.x_values, s.values, s.sizes):
                ser.add_data_point(x, y, z)
        return data
    if not item.categories or any(len(s.values) != len(item.categories) for s in item.series):
        return None
    data = CategoryChartData(number_format=item.series[0].number_format or "General")
    data.categories = item.categories
    for s in item.series:
        data.add_series(s.name, s.values, s.number_format)
    return data


def _numbering_positions(deck):
    """1-based position of each numbered paragraph within its original list."""
    positions = {}
    for s in deck.slides:
        counters = {}
        last_box = None
        for t in s.texts:
            if t.box != last_box:
                counters, last_box = {}, t.box
            for lvl in [k for k in counters if k > t.level]:
                del counters[lvl]
            if t.bullet == "number":
                counters[t.level] = counters.get(t.level, 0) + 1
                positions[t.id] = counters[t.level]
            else:
                counters.pop(t.level, None)
    return positions


def _balanced_split(body):
    """Index of the top-level paragraph that best halves a list, or None."""
    top = min(b.level for b in body)
    weights = [max(len(b.text), 20) for b in body]
    total = sum(weights)
    best, best_diff = None, None
    running = 0
    for i, b in enumerate(body):
        if i and b.level == top:
            diff = abs(total - 2 * running)
            if best_diff is None or diff < best_diff:
                best, best_diff = i, diff
        running += weights[i]
    return best
