"""Amendment 1C (locked objects) and 1E (images) in the extract step."""

import hashlib
import io

import pytest
from pptx import Presentation

from lockedbuilder import build_images_deck, build_locked_deck
from polisher.extract import extract_deck


@pytest.fixture(scope="module")
def locked_path(tmp_path_factory):
    return build_locked_deck(tmp_path_factory.mktemp("lk") / "locked.pptx")


@pytest.fixture(scope="module")
def locked(locked_path):
    return extract_deck(locked_path)


@pytest.fixture(scope="module")
def images(tmp_path_factory):
    return extract_deck(build_images_deck(tmp_path_factory.mktemp("im") / "images.pptx"))


EXPECTED = ["video", "audio", "ole", "ole", "equation", "model3d", "smartart", "ink", "group"]


def test_each_locked_type_is_detected(locked):
    found = [[x.lock_type for x in s.locked] for s in locked.slides[1:]]
    assert found == [[t] for t in EXPECTED]


def test_nothing_is_unsupported_any_more(locked):
    assert locked.unsupported == []


def test_locked_items_record_shape_id_and_position(locked, locked_path):
    prs = Presentation(locked_path)
    video = locked.slides[1].locked[0]
    pic = next(sh for sh in prs.slides[1].shapes if sh.shape_type is not None and "movie" in sh.name)
    assert video.shape_id == pic.shape_id
    assert video.id == "s2-l1"
    x, y, w, h = video.pos
    assert 0 < x < 1 and 0 < y < 1 and w > 0 and h > 0
    assert video.aspect == pytest.approx(4 / 3, rel=0.01)


def test_part_hashes_cover_every_related_part(locked, locked_path):
    prs = Presentation(locked_path)
    slide = prs.slides[1]
    blobs = {r.target_part.blob for r in slide.part.rels.values()
             if not r.is_external and r.reltype.rsplit("/", 1)[-1] in ("media", "video", "image")}
    assert sorted(locked.slides[1].locked[0].part_hashes) == sorted({hashlib.sha256(b).hexdigest() for b in blobs})


def test_smartart_hashes_all_four_diagram_parts(locked):
    assert len(locked.slides[7].locked[0].part_hashes) == 4


def test_linked_object_is_flagged(locked):
    ole = locked.slides[4].locked[0]
    assert ole.links == ["file:///C:/Data/Sales%202026.xlsx"]
    assert [x.id for x in locked.linked_objects()] == [ole.id]


def test_equations_are_counted_and_not_read_as_text(locked):
    s = locked.slides[5]
    assert s.locked[0].equations == 1
    assert [t.text for t in s.texts] == ["The growth formula", "Where r is the monthly growth rate."]


def test_smartart_text_is_not_extracted_as_paragraphs(locked):
    assert [t.text for t in locked.slides[7].texts] == ["Our process"]


def test_text_inside_a_locked_group_stays_with_the_group(locked):
    s = locked.slides[9]
    assert [t.text for t in s.texts] == ["Store opening"]
    assert s.locked[0].lock_type == "group"


def test_locked_summary_for_the_quote(locked):
    summary = locked.locked_summary()
    assert summary["slides"] == 9
    assert summary["counts"]["ole"] == 2 and summary["counts"]["equation"] == 1
    assert summary["linked"] == 1
    assert summary["has_media"] is True


# --- images (Amendment 1E) ---------------------------------------------------

def test_master_and_layout_images_are_recorded(images):
    containers = sorted(m.container for m in images.master_images)
    assert containers == ["layout:2", "master:1"]
    assert all(len(m.sha256) == 64 for m in images.master_images)


def test_crop_is_recorded_exactly(images):
    (img,) = images.slides[1].images
    assert img.crop_raw == {"l": 15000, "t": 10000, "r": 5000, "b": 0}


def test_alt_text_and_title_are_recorded(images):
    (img,) = images.slides[1].images
    assert (img.alt_text, img.title) == ("Warehouse exterior at dawn", "Warehouse")


def test_picture_background_is_an_item(images):
    s = images.slides[2]
    (bg,) = s.backgrounds
    assert bg.id == "s3-bg1" and len(bg.sha256) == 64
    assert bg.aspect == pytest.approx(800 / 600, rel=0.01)


def test_svg_pictures_keep_both_parts(images):
    icons = images.slides[3].images
    assert len(icons) == 3
    assert all(i.svg_sha256 and i.svg_sha256 != i.sha256 for i in icons)


def test_items_record_their_shape_ids(images):
    s = images.slides[1]
    assert s.images[0].shape_id > 0 and all(t.shape_id > 0 for t in s.texts)
