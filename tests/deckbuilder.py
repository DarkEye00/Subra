"""Builders for synthetic test decks.

Real customer decks are the real test (see PROGRESS.md), but these cover the
structures the extractor has to handle: placeholders, bullets with levels,
runs, tables with merges, charts, cropped images, groups, notes, hidden slides.
"""

import io
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt


def png_bytes(width=400, height=300, color=(200, 30, 30)):
    buf = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buf, format="PNG")
    return buf.getvalue()


def build_rich_deck(path):
    """A 7-slide deck exercising every content type we extract."""
    prs = Presentation()

    # 1. Title slide: title + subtitle placeholders
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Quarterly Business Review"
    s.placeholders[1].text = "Nairobi office, Q3 2026"
    s.notes_slide.notes_text_frame.text = "Welcome everyone.\nKeep intro under 2 minutes."

    # 2. Bullets with levels and mixed runs
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Highlights"
    tf = s.placeholders[1].text_frame
    tf.text = "Revenue grew 45% to KES 12.5m"
    p = tf.add_paragraph()
    p.text = "Driven by M-Pesa collections"
    p.level = 1
    p = tf.add_paragraph()
    r1 = p.add_run()
    r1.text = "Churn: "
    r1.font.bold = True
    r2 = p.add_run()
    r2.text = "down to 3.2%"
    r2.font.italic = True
    s.notes_slide.notes_text_frame.text = "Mention the Mombasa launch."

    # 3. Table with a merged header cell
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Regional sales"
    rows, cols = 4, 3
    tbl = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.5), Inches(9), Inches(3)).table
    data = [
        ["Region", "2025", "2026"],
        ["Nairobi", "1,200", "1,850"],
        ["Mombasa", "640", "910"],
        ["Kisumu", "300", "415"],
    ]
    for r in range(rows):
        for c in range(cols):
            tbl.cell(r, c).text = data[r][c]
    tbl.cell(0, 1).merge(tbl.cell(0, 2))
    tbl.cell(0, 1).text = "Units sold"

    # 4. Bar chart with two series
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Growth by quarter"
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3", "Q4"]
    cd.add_series("2025", (10.5, 12.0, 14.25, 15.0))
    cd.add_series("2026", (13.0, 16.5, 21.0, None))
    gf = s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(5), cd
    )
    gf.chart.has_title = True
    gf.chart.chart_title.text_frame.text = "Revenue (KES m)"

    # 5. Cropped image + caption text box
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic = s.shapes.add_picture(io.BytesIO(png_bytes(400, 300)), Inches(1), Inches(1), Inches(4))
    pic.crop_left = 0.1
    pic.crop_right = 0.1
    tb = s.shapes.add_textbox(Inches(6), Inches(1), Inches(3), Inches(1))
    tb.text_frame.text = "Our new Westlands branch"

    # 6. Group shape containing text, plus a scatter chart
    s = prs.slides.add_slide(prs.slide_layouts[6])
    grp = s.shapes.add_group_shape()
    a = grp.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    a.text_frame.text = "Grouped note A"
    b = grp.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    b.text_frame.text = "Grouped note B"
    xy = XyChartData()
    ser = xy.add_series("Price vs volume")
    ser.add_data_point(1.5, 200)
    ser.add_data_point(2.0, 180)
    ser.add_data_point(3.25, 120)
    s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(5), Inches(1), Inches(4), Inches(4), xy)

    # 7. Hidden closing slide
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Thank you"
    s.placeholders[1].text = "questions@example.co.ke"
    s._element.set("show", "0")

    prs.save(path)
    return path


def build_simple_deck(path, n_slides=3):
    prs = Presentation()
    for i in range(n_slides):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Slide {i + 1}"
        s.placeholders[1].text = f"Point {i + 1}"
    prs.save(path)
    return path


def add_zip_member(src, dst, name, data):
    """Copy a zip package and add one extra member (e.g. vbaProject.bin)."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr(name, data)
    return dst


def cfb_bytes(with_encrypted_package):
    """Minimal bytes that look like an OLE compound file (legacy .ppt or encrypted OOXML)."""
    header = bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 504
    body = "EncryptedPackage".encode("utf-16-le") if with_encrypted_package else b"PowerPoint Document"
    return header + body + b"\x00" * 1024
