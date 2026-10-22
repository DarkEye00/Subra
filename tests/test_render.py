import copy

import pytest
from pptx import Presentation
from pptx.util import Inches

from deckbuilder import build_rich_deck
from polisher.extract import extract_deck
from polisher.heuristic import plan_deck
from polisher.plan import parse_plan
from polisher.render.renderer import render
from polisher.render.themes import THEMES, get_theme
from polisher.verify import verify


@pytest.fixture(scope="module")
def rich_path(tmp_path_factory):
    return build_rich_deck(tmp_path_factory.mktemp("r") / "rich.pptx")


@pytest.fixture(scope="module")
def rich(rich_path):
    deck = extract_deck(rich_path)
    deck.source = rich_path
    return deck


def render_to(tmp_path, deck, plan, theme="navy", name="out.pptx", source=None):
    out = tmp_path / name
    report = render(plan, deck, get_theme(theme), out, source or deck.source)
    return out, report


@pytest.mark.parametrize("theme", sorted(THEMES))
def test_round_trip_keeps_all_content(tmp_path, rich, theme):
    out, _ = render_to(tmp_path, rich, plan_deck(rich), theme)
    result = verify(rich, extract_deck(out))
    assert result.passed, result.problems


def test_output_keeps_the_original_slide_size(tmp_path, rich, rich_path):
    out, _ = render_to(tmp_path, rich, plan_deck(rich))
    before, after = Presentation(rich_path), Presentation(out)
    assert (after.slide_width, after.slide_height) == (before.slide_width, before.slide_height)


def test_theme_fonts_are_written_into_the_office_theme(tmp_path, rich):
    out, _ = render_to(tmp_path, rich, plan_deck(rich), "navy")
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    theme_xml = Presentation(out).slide_master.part.part_related_by(RT.THEME).blob.decode()
    assert '<a:majorFont><a:latin typeface="Cambria"' in theme_xml
    assert '<a:minorFont><a:latin typeface="Calibri"' in theme_xml


def test_content_titles_use_the_title_placeholder(tmp_path, rich):
    out, _ = render_to(tmp_path, rich, plan_deck(rich))
    prs = Presentation(out)
    assert prs.slides[1].shapes.title.text == "Highlights"


def test_tables_and_charts_are_native_and_editable(tmp_path, rich):
    out, _ = render_to(tmp_path, rich, plan_deck(rich))
    prs = Presentation(out)
    assert any(getattr(s, "has_table", False) and s.has_table for s in prs.slides[2].shapes)
    assert any(getattr(s, "has_chart", False) and s.has_chart for s in prs.slides[3].shapes)


def test_hidden_slides_stay_hidden(tmp_path, rich):
    out, _ = render_to(tmp_path, rich, plan_deck(rich))
    assert extract_deck(out).slides[-1].hidden is True


def test_split_slide_round_trips(tmp_path, rich):
    plan = plan_deck(rich)
    data = plan.to_dict()
    s2 = data["slides"][1]
    s2["slots"]["body"] = ["s2-t2", "s2-t3"]
    data["slides"].insert(2, {"source_slide": 2, "layout": "bullets", "slots": {"body": ["s2-t4"]}})
    out, _ = render_to(tmp_path, rich, parse_plan(data))
    rendered = extract_deck(out)
    assert len(rendered.slides) == len(rich.slides) + 1
    assert verify(rich, rendered).passed


def test_stat_highlight_is_big_but_text_is_unchanged(tmp_path, rich):
    data = plan_deck(rich).to_dict()
    data["slides"][1] = {
        "source_slide": 2,
        "layout": "stats",
        "slots": {
            "title": ["s2-t1"],
            "stats": [{"ids": ["s2-t2"], "highlight": "Revenue grew 45%"}, {"ids": ["s2-t4"]}],
            "body": ["s2-t3"],
        },
    }
    out, _ = render_to(tmp_path, rich, parse_plan(data))
    assert verify(rich, extract_deck(out)).passed
    prs = Presentation(out)
    card = next(sh for sh in prs.slides[1].shapes if sh.has_text_frame and sh.text_frame.text.startswith("Revenue"))
    runs = [r for r in card.text_frame.paragraphs[0].runs]
    assert runs[0].text == "Revenue grew 45%"
    assert runs[0].font.size > runs[-1].font.size


def test_every_layout_round_trips(tmp_path, rich):
    data = {
        "slides": [
            {"source_slide": 1, "layout": "section", "slots": {"title": ["s1-t1"], "subtitle": ["s1-t2"]}},
            {"source_slide": 2, "layout": "two_column",
             "slots": {"title": ["s2-t1"], "left_title": ["s2-t2"], "left": ["s2-t3"], "right": ["s2-t4"]}},
            {"source_slide": 3, "layout": "table", "slots": {"title": ["s3-t1"], "visuals": ["s3-tb1"]}},
            {"source_slide": 4, "layout": "image_text", "slots": {"title": ["s4-t1"], "visuals": ["s4-c1"]}},
            {"source_slide": 5, "layout": "image_text", "slots": {"visuals": ["s5-i1"], "body": ["s5-t1"]}},
            {"source_slide": 6, "layout": "table", "slots": {"title": ["s6-t1"], "visuals": ["s6-c1"], "body": ["s6-t2"]}},
            {"source_slide": 7, "layout": "title", "slots": {"title": ["s7-t1"], "subtitle": ["s7-t2"]}},
        ]
    }
    out, _ = render_to(tmp_path, rich, parse_plan(data))
    result = verify(rich, extract_deck(out))
    assert result.passed, result.problems


def test_kept_chart_round_trips(tmp_path, rich):
    deck = copy.deepcopy(rich)
    deck.slides[3].charts[0].rebuild = "keep"
    out, _ = render_to(tmp_path, deck, plan_deck(deck))
    result = verify(deck, extract_deck(out))
    assert result.passed, result.problems


def test_overflowing_slide_is_flagged(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Too much"
    tf = s.placeholders[1].text_frame
    tf.text = "word " * 400
    for _ in range(30):
        tf.add_paragraph().text = "another long paragraph with many words in it " * 3
    prs.save(tmp_path / "big.pptx")
    deck = extract_deck(tmp_path / "big.pptx")
    deck.source = tmp_path / "big.pptx"
    ids = [t.id for t in deck.slides[0].texts]
    plan = parse_plan({"slides": [{"source_slide": 1, "layout": "bullets",
                                   "slots": {"title": ids[:1], "body": ids[1:]}}]})
    out, report = render_to(tmp_path, deck, plan)
    assert report.overflow_slides == [1]
    assert verify(deck, extract_deck(out)).passed


def test_footer_text_is_kept(tmp_path):
    from pptx.enum.shapes import PP_PLACEHOLDER

    prs = Presentation()
    layout = prs.slide_layouts[1]
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = "Hello"
    s.placeholders[1].text = "Body"
    ph = next(p for p in layout.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.FOOTER)
    s.shapes.clone_placeholder(ph)
    next(p for p in s.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.FOOTER).text = "Confidential"
    prs.save(tmp_path / "f.pptx")
    deck = extract_deck(tmp_path / "f.pptx")
    deck.source = tmp_path / "f.pptx"
    out, _ = render_to(tmp_path, deck, plan_deck(deck))
    assert verify(deck, extract_deck(out)).passed


def test_long_list_of_short_lines_flows_into_two_columns(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Cover"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Operations"
    tf = s.placeholders[1].text_frame
    tf.text = "Group 0"
    for g in range(5):
        if g:
            tf.add_paragraph().text = f"Group {g}"
        for i in range(2):
            p = tf.add_paragraph()
            p.text = f"Detail {g}.{i} about the work"
            p.level = 1
    prs.save(tmp_path / "d.pptx")
    deck = extract_deck(tmp_path / "d.pptx")
    deck.source = tmp_path / "d.pptx"
    out, _ = render_to(tmp_path, deck, plan_deck(deck))
    body_boxes = [sh for sh in Presentation(out).slides[1].shapes
                  if sh.has_text_frame and sh.text_frame.text.startswith("Group")]
    assert len(body_boxes) == 2
    assert verify(deck, extract_deck(out)).passed


# --- Task 2: 12pt floor and named shapes ------------------------------------

def _run_sizes(prs, in_tables):
    """Explicit run sizes (pt) on every slide, either inside tables or outside them."""
    sizes = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                if in_tables:
                    for cell in sh.table.iter_cells():
                        sizes += [r.font.size.pt for p in cell.text_frame.paragraphs for r in p.runs if r.font.size]
            elif sh.has_text_frame and not in_tables:
                sizes += [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
    return sizes


@pytest.mark.parametrize("theme", sorted(THEMES))
def test_no_text_below_twelve_points_on_the_business_deck(tmp_path, theme):
    from deckbuilder import build_business_deck

    src = build_business_deck(tmp_path / "b.pptx")
    deck = extract_deck(src)
    out, _ = render_to(tmp_path, deck, plan_deck(deck), theme, source=src)
    sizes = _run_sizes(Presentation(out), in_tables=False)
    assert sizes and min(sizes) >= 12
    assert min(_run_sizes(Presentation(out), in_tables=True)) >= 12


def test_long_table_floors_at_ten_points_and_is_flagged(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Every branch"
    tbl = s.shapes.add_table(30, 4, Inches(0.5), Inches(1.5), Inches(9), Inches(5)).table
    for r in range(30):
        for c in range(4):
            tbl.cell(r, c).text = f"Branch {r}" if c == 0 else f"{r * 10 + c},{c}00"
    prs.save(tmp_path / "t.pptx")
    deck = extract_deck(tmp_path / "t.pptx")
    out, report = render_to(tmp_path, deck, plan_deck(deck), source=tmp_path / "t.pptx")
    sizes = _run_sizes(Presentation(out), in_tables=True)
    assert 10 <= min(sizes) < 12
    assert report.overflow_slides == [1]
    assert report.slides[0]["min_pt"] < 12
    assert verify(deck, extract_deck(out)).passed


def test_drawn_shapes_are_named(tmp_path, rich):
    out, _ = render_to(tmp_path, rich, plan_deck(rich))
    names = {sh.name for slide in Presentation(out).slides for sh in slide.shapes}
    assert {"SUBRA decor", "SUBRA text", "SUBRA table", "SUBRA chart"} <= names, names
    # the chart slide with a text panel draws a card
    assert "SUBRA card" in names
