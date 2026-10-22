"""Capacity limits of the layout library (17 Oct build plan: limits are enforced in code)."""

import pytest
from pptx import Presentation

from deckbuilder import build_business_deck
from polisher.extract import extract_deck
from polisher.heuristic import plan_deck
from polisher.layouts import LAYOUTS
from polisher.plan import parse_plan, repair_plan, validate_plan


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    return extract_deck(build_business_deck(tmp_path_factory.mktemp("l") / "business.pptx"))


@pytest.fixture(scope="module")
def deck8(tmp_path_factory):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Cover"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Agenda"
    tf = s.placeholders[1].text_frame
    tf.text = "Item 1"
    for i in range(2, 9):
        tf.add_paragraph().text = f"Item {i}"
    path = tmp_path_factory.mktemp("l8") / "agenda8.pptx"
    prs.save(path)
    return extract_deck(path)


def assert_error(errors, *needles):
    assert any(all(n in e for n in needles) for e in errors), errors


def test_library_has_twelve_layouts():
    assert sorted(LAYOUTS) == ["agenda", "bullets", "chart", "image_full", "image_text", "quote", "section",
                               "stats", "table", "timeline", "title", "two_column"]


@pytest.mark.parametrize("layout,slot,lo,hi", [
    ("agenda", "items", 3, 7), ("timeline", "steps", 3, 6), ("image_full", "visuals", 1, 1), ("stats", "stats", 1, 4),
])
def test_capacity_limits_are_declared(layout, slot, lo, hi):
    s = LAYOUTS[layout].slot(slot)
    assert (s.min_items, s.max_items) == (lo, hi)


def test_agenda_with_two_items_is_invalid_and_repairs_to_bullets(deck):
    data = plan_deck(deck).to_dict()
    data["slides"][1] = {"source_slide": 2, "layout": "agenda", "slots": {"title": ["s2-t1"], "items": ["s2-t2", "s2-t3"]}}
    data["slides"].insert(2, {"source_slide": 2, "layout": "bullets", "slots": {"body": ["s2-t4", "s2-t5", "s2-t6"]}})
    plan = parse_plan(data)
    assert_error(validate_plan(plan, deck), "items", "at least 3")
    repaired, _ = repair_plan(plan, deck)
    assert validate_plan(repaired, deck) == []
    assert repaired.slides[1].layout == "bullets"


def test_agenda_with_eight_items_splits(deck8):
    ids = [t.id for t in deck8.slides[1].texts if t.role != "title"]
    assert len(ids) == 8
    plan = parse_plan({"slides": [
        {"source_slide": 1, "layout": "title", "slots": {"title": ["s1-t1"]}},
        {"source_slide": 2, "layout": "agenda", "slots": {"title": ["s2-t1"], "items": ids}},
    ]})
    assert_error(validate_plan(plan, deck8), "items", "at most 7")
    repaired, _ = repair_plan(plan, deck8)
    assert validate_plan(repaired, deck8) == []
    assert [p.layout for p in repaired.slides if p.source_slide == 2] == ["agenda", "bullets"]
