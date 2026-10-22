from pptx import Presentation
from pptx.util import Inches

from deckbuilder import build_rich_deck
from polisher.extract import extract_deck
from polisher.heuristic import plan_deck
from polisher.plan import validate_plan


def deck_from(tmp_path, build):
    prs = Presentation()
    build(prs)
    prs.save(tmp_path / "d.pptx")
    return extract_deck(tmp_path / "d.pptx")


def bullets_slide(prs, title, paragraphs):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame
    tf.text = paragraphs[0]
    for p in paragraphs[1:]:
        tf.add_paragraph().text = p


def test_heuristic_plan_for_rich_deck_is_valid(tmp_path):
    deck = extract_deck(build_rich_deck(tmp_path / "rich.pptx"))
    plan = plan_deck(deck)
    assert validate_plan(plan, deck) == []
    assert [p.layout for p in plan.slides] == ["title", "bullets", "table", "chart", "image_text", "chart", "title"]


def test_number_led_paragraphs_become_stats_with_highlights(tmp_path):
    def build(prs):
        bullets_slide(prs, "Cover", ["x"])
        bullets_slide(prs, "Results", ["45% revenue growth", "KES 12.5m raised", "3,000"])
        bullets_slide(prs, "End", ["a longer closing paragraph that is not short enough to be a quote at all" * 2])

    deck = deck_from(tmp_path, build)
    plan = plan_deck(deck)
    stats = plan.slides[1]
    assert stats.layout == "stats"
    assert [(e.ids, e.highlight) for e in stats.slots["stats"]] == [
        (["s2-t2"], "45%"),
        (["s2-t3"], "KES 12.5m"),
        (["s2-t4"], None),
    ]
    assert validate_plan(plan, deck) == []


def test_overlong_bullets_are_split_across_slides(tmp_path):
    def build(prs):
        bullets_slide(prs, "Cover", ["x"])
        bullets_slide(prs, "Everything", [f"Point {i}: " + "detail " * 25 for i in range(12)])

    deck = deck_from(tmp_path, build)
    plan = plan_deck(deck)
    parts = [p for p in plan.slides if p.source_slide == 2]
    assert len(parts) > 1
    assert "title" in parts[0].slots and all("title" not in p.slots for p in parts[1:])
    assert validate_plan(plan, deck) == []


def test_slide_without_content_is_blank(tmp_path):
    def build(prs):
        bullets_slide(prs, "Cover", ["x"])
        prs.slides.add_slide(prs.slide_layouts[6])
        bullets_slide(prs, "After", ["a long enough paragraph so this is not treated as a closing quote slide" * 2])

    deck = deck_from(tmp_path, build)
    plan = plan_deck(deck)
    assert plan.slides[1].layout == "blank"
    assert validate_plan(plan, deck) == []


def test_agenda_slide_gets_the_agenda_layout(tmp_path):
    from deckbuilder import build_business_deck

    deck = extract_deck(build_business_deck(tmp_path / "b.pptx"))
    plan = plan_deck(deck)
    assert plan.slides[1].layout == "agenda"
    assert len(plan.slides[1].slots["items"]) == 5
    assert validate_plan(plan, deck) == []


def test_single_large_picture_gets_image_full(tmp_path):
    import io

    from deckbuilder import photo_bytes

    def build(prs):
        bullets_slide(prs, "Cover", ["x"])
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = "Our new hub"
        s.shapes.add_picture(io.BytesIO(photo_bytes(seed=2)), Inches(0.5), Inches(1.5), Inches(9))
        bullets_slide(prs, "After", ["a long enough paragraph so this is not treated as a closing quote slide" * 2])

    deck = deck_from(tmp_path, build)
    plan = plan_deck(deck)
    assert plan.slides[1].layout == "image_full"
    assert validate_plan(plan, deck) == []


# --- fallback plans (the plainest safe layout for a slide that keeps failing) ---

def _items_of(deck, n):
    from polisher.plan import _plannable

    items = _plannable(deck)
    return [items[i] for i in deck.slides[n - 1].order if i in items]


def test_fallback_plan_splits_text_in_halves_then_quarters(tmp_path):
    from polisher.heuristic import fallback_plan, plan_items
    from polisher.plan import Plan

    def build(prs):
        bullets_slide(prs, "Cover", ["x"])
        bullets_slide(prs, "Everything", [f"Point {i}: " + "detail " * 25 for i in range(12)])

    deck = deck_from(tmp_path, build)
    items = _items_of(deck, 2)
    default, one, two = plan_items(2, items), fallback_plan(2, items, level=1), fallback_plan(2, items, level=2)
    assert len(default) < len(one) < len(two)
    assert all(p.layout == "bullets" for p in one + two)
    assert "title" in one[0].slots and all("title" not in p.slots for p in one[1:])
    cover = plan_items(1, _items_of(deck, 1), is_first_slide=True)
    assert validate_plan(Plan(cover + two), deck) == []


def test_fallback_plan_puts_visuals_first_with_the_title(tmp_path):
    from polisher.heuristic import fallback_plan

    deck = extract_deck(build_rich_deck(tmp_path / "rich.pptx"))
    parts = fallback_plan(6, _items_of(deck, 6))  # two grouped notes (the first reads as the title) and a chart
    assert parts[0].layout == "chart" and parts[0].slots == {"title": ["s6-t1"], "visuals": ["s6-c1"]}
    assert parts[1].layout == "bullets" and parts[1].slots == {"body": ["s6-t2"]}
    parts = fallback_plan(3, _items_of(deck, 3))  # a title and a table
    assert parts == [parts[0]] and parts[0].layout == "table" and parts[0].slots["title"] == ["s3-t1"]
