"""Deterministic layout checks on a rendered deck (17 Oct build plan, stage 6)."""

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from deckbuilder import build_business_deck, build_rich_deck
from polisher.extract import extract_deck
from polisher.heuristic import plan_deck
from polisher.plan import parse_plan
from polisher.qa import QAReport, check_deck, contrast_ratio
from polisher.render.renderer import RenderReport, render
from polisher.render.themes import THEMES, get_theme

NAVY = get_theme("navy")


def _report(n=1, **entry):
    base = {"index": 1, "source_slide": 1, "layout": "bullets", "overflow": False, "min_pt": 18.0}
    return RenderReport(slides=[{**base, **entry, "index": i + 1, "source_slide": i + 1} for i in range(n)])


def _blank_deck(tmp_path, name="d.pptx"):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return prs, slide, tmp_path / name


def _textbox(slide, text, x, y, w, h, size=18, color="1F2A37", name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return tb


def _rect(slide, x, y, w, h, color, name):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = name
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(color)
    return sp


def _kinds(qa):
    return [f.kind for f in qa.slides[0].faults]


def test_contrast_ratio_matches_wcag_reference():
    assert contrast_ratio("000000", "FFFFFF") == pytest.approx(21.0)
    assert contrast_ratio("777777", "FFFFFF") == pytest.approx(4.48, abs=0.01)
    assert contrast_ratio("FFFFFF", "000000") == contrast_ratio("000000", "FFFFFF")


@pytest.mark.parametrize("theme", sorted(THEMES))
@pytest.mark.parametrize("builder", [build_rich_deck, build_business_deck])
def test_clean_render_has_no_faults(tmp_path, theme, builder):
    src = builder(tmp_path / "in.pptx")
    deck = extract_deck(src)
    out = tmp_path / "out.pptx"
    report = render(plan_deck(deck), deck, get_theme(theme), out, src)
    qa = check_deck(out, report, get_theme(theme))
    assert qa.flagged_slides == [], [(s.index, f.kind, f.detail) for s in qa.faulted for f in s.faults]
    assert len(qa.slides) == len(report.slides)


def test_overlapping_text_boxes_are_reported(tmp_path):
    prs, slide, path = _blank_deck(tmp_path)
    _textbox(slide, "Revenue grew 45%", 1, 1, 4, 1)
    _textbox(slide, "Margin held at 38%", 2, 1.3, 4, 1)
    prs.save(path)
    qa = check_deck(path, _report(), NAVY)
    assert _kinds(qa) == ["overlap"]
    fault = qa.slides[0].faults[0]
    assert "Revenue grew 45%" in fault.detail and "Margin held" in fault.detail and fault.repairable


def test_decoration_and_cards_do_not_count_as_overlap(tmp_path):
    prs, slide, path = _blank_deck(tmp_path)
    _rect(slide, 0.5, 0.5, 5, 3, "F1F4F8", "SUBRA card")
    _rect(slide, 0, 0, 0.12, 7.5, "0B1F3A", "SUBRA decor")
    _textbox(slide, "Inside the card", 1, 1, 4, 1)
    prs.save(path)
    assert _kinds(check_deck(path, _report(), NAVY)) == []


def test_overlay_band_may_cover_a_picture(tmp_path):
    import io

    from deckbuilder import png_bytes

    prs, slide, path = _blank_deck(tmp_path)
    slide.shapes.add_picture(io.BytesIO(png_bytes()), Inches(0), Inches(0), Inches(10), Inches(6.2))
    _rect(slide, 0, 6.2, 10, 1.3, "0B1F3A", "SUBRA overlay")
    _textbox(slide, "The Mombasa hub", 0.75, 6.4, 8.5, 0.9, color="FFFFFF")
    prs.save(path)
    assert _kinds(check_deck(path, _report(layout="image_full"), NAVY)) == []


def test_tiny_text_is_reported(tmp_path):
    prs, slide, path = _blank_deck(tmp_path)
    _textbox(slide, "Fine print", 1, 1, 4, 1, size=9, name="SUBRA text")
    prs.save(path)
    qa = check_deck(path, _report(), NAVY)
    assert _kinds(qa) == ["tiny_text"]
    assert "9pt" in qa.slides[0].faults[0].detail and qa.slides[0].faults[0].repairable


def test_tiny_text_inside_a_table_is_not_repairable(tmp_path):
    prs, slide, path = _blank_deck(tmp_path)
    frame = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2))
    frame.name = "SUBRA table"
    tbl = frame.table
    for r in range(2):
        for c in range(2):
            run = tbl.cell(r, c).text_frame.paragraphs[0].add_run()
            run.text = f"cell {r}{c}"
            run.font.size = Pt(10)
    prs.save(path)
    qa = check_deck(path, _report(layout="table"), NAVY)
    assert _kinds(qa) == ["tiny_text"] and qa.slides[0].faults[0].repairable is False


def test_shape_off_the_slide_is_reported(tmp_path):
    prs, slide, path = _blank_deck(tmp_path)
    _textbox(slide, "Too low", 1, 7, 4, 1.5)  # bottom at 8.5in on a 7.5in slide
    prs.save(path)
    qa = check_deck(path, _report(), NAVY)
    assert _kinds(qa) == ["off_slide"] and "Too low" in qa.slides[0].faults[0].detail


def test_low_contrast_is_reported(tmp_path):
    prs, slide, path = _blank_deck(tmp_path)
    _textbox(slide, "Faint", 1, 1, 4, 1, color="BBBBBB", name="SUBRA text")
    prs.save(path)
    qa = check_deck(path, _report(), NAVY)
    assert _kinds(qa) == ["low_contrast"] and qa.slides[0].faults[0].repairable is False


def test_contrast_uses_the_card_under_the_text(tmp_path):
    prs, slide, path = _blank_deck(tmp_path)
    _rect(slide, 0.5, 0.5, 5, 3, "0B1F3A", "SUBRA card")
    _textbox(slide, "White on navy", 1, 1, 4, 1, color="FFFFFF", name="SUBRA text")
    prs.save(path)
    assert _kinds(check_deck(path, _report(), NAVY)) == []


def test_empty_placeholder_is_reported(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Only a title"
    path = tmp_path / "p.pptx"
    prs.save(path)
    qa = check_deck(path, _report(), NAVY)
    assert _kinds(qa) == ["empty_placeholder"]


def test_overflow_flag_from_the_render_report_is_a_fault(tmp_path):
    prs, slide, path = _blank_deck(tmp_path)
    _textbox(slide, "Fits fine", 1, 1, 4, 1)
    prs.save(path)
    qa = check_deck(path, _report(overflow=True, min_pt=12.0), NAVY)
    assert _kinds(qa) == ["overflow"]
    assert qa.faults_by_source() == {1: ["Planned slide 1 (layout bullets): text does not fit its box even at 12pt"]}
    assert qa.repairable_sources == [1] and qa.flagged_slides == [1]


def test_table_overflow_detail_names_the_floor(tmp_path):
    prs, slide, path = _blank_deck(tmp_path)
    _textbox(slide, "x", 1, 1, 4, 1)
    prs.save(path)
    qa = check_deck(path, _report(overflow=True, min_pt=10.0, layout="table"), NAVY)
    assert qa.slides[0].faults[0].detail == "table text is below 12pt"


def test_mark_sets_repaired_and_fallback_flags():
    qa = QAReport(slides=[])
    from polisher.qa import SlideQA

    qa.slides = [SlideQA(1, 1, "title"), SlideQA(2, 2, "bullets"), SlideQA(3, 2, "bullets"), SlideQA(4, 3, "chart")]
    qa.mark(repaired_sources=[2], fallback_sources=[3])
    assert [s.repaired for s in qa.slides] == [False, True, True, False]
    assert [s.fallback for s in qa.slides] == [False, False, False, True]
    d = qa.to_dict()
    assert d["slides"][1]["repaired"] is True and d["flagged_slides"] == []


# --- review fixes: text SUBRA kept verbatim is not its own fault ------------

def test_kept_footer_text_is_not_checked_for_size_or_contrast(tmp_path):
    from pptx.enum.shapes import PP_PLACEHOLDER

    prs = Presentation()
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Only a title"
    slide.placeholders[1].text = "Body"
    lf = next(p for p in layout.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.FOOTER)
    slide.shapes.clone_placeholder(lf)
    footer = next(p for p in slide.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.FOOTER)
    footer.left, footer.top, footer.width, footer.height = Inches(0.5), Inches(7.0), Inches(9), Inches(0.4)
    run = footer.text_frame.paragraphs[0].add_run()
    run.text = "Confidential"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string("CCCCCC")
    path = tmp_path / "f.pptx"
    prs.save(path)
    assert _kinds(check_deck(path, _report(), NAVY)) == []


def test_text_inside_a_kept_group_is_not_checked_for_size(tmp_path):
    prs, slide, path = _blank_deck(tmp_path)
    grp = slide.shapes.add_group_shape()
    tb = grp.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "A caption the customer wrote at 10pt"
    run.font.size = Pt(10)
    prs.save(path)
    assert _kinds(check_deck(path, _report(), NAVY)) == []


def test_tiny_text_in_the_renderers_own_boxes_is_still_reported(tmp_path):
    prs, slide, path = _blank_deck(tmp_path)
    _textbox(slide, "Fine print", 1, 1, 4, 1, size=9, name="SUBRA text")
    prs.save(path)
    assert _kinds(check_deck(path, _report(), NAVY)) == ["tiny_text"]


def test_inherited_footer_position_is_seen_by_the_overlap_check(tmp_path):
    from pptx.enum.shapes import PP_PLACEHOLDER

    prs = Presentation()
    layout = prs.slide_layouts[5]  # title only, with footer placeholders
    lf = next(p for p in layout.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.FOOTER)
    lf.left, lf.top, lf.width, lf.height = Inches(0.5), Inches(3.0), Inches(9), Inches(0.4)
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Title"
    slide.shapes.clone_placeholder(lf)
    next(p for p in slide.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.FOOTER).text = "Confidential"
    _textbox(slide, "Body text over the footer", 1, 2.8, 4, 1, name="SUBRA text")
    path = tmp_path / "f.pptx"
    prs.save(path)
    assert _kinds(check_deck(path, _report(), NAVY)) == ["overlap"]
