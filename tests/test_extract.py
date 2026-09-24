import hashlib
import io
import json

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from deckbuilder import build_rich_deck, png_bytes
from polisher.extract import extract_deck, parse_smartart_data


@pytest.fixture(scope="module")
def rich(tmp_path_factory):
    path = build_rich_deck(tmp_path_factory.mktemp("decks") / "rich.pptx")
    return extract_deck(path)


def texts(slide):
    return [t.text for t in slide.texts]


# --- identity -------------------------------------------------------------

def test_every_item_has_a_unique_id(rich):
    ids = [item.id for s in rich.slides for item in s.items()]
    assert len(ids) == len(set(ids))
    assert all(i.startswith(f"s{s.number}-") for s in rich.slides for i in [x.id for x in s.items()])


def test_ids_are_stable_across_extractions(tmp_path):
    path = build_rich_deck(tmp_path / "d.pptx")
    a, b = extract_deck(path), extract_deck(path)
    assert a.to_dict() == b.to_dict()


def test_to_dict_is_json_serializable_and_excludes_image_bytes(rich):
    blob = json.dumps(rich.to_dict())
    assert "blob" not in blob
    assert len(blob) < 50_000


# --- text -----------------------------------------------------------------

def test_title_slide_title_and_subtitle_roles(rich):
    s1 = rich.slides[0]
    assert [(t.role, t.text) for t in s1.texts] == [
        ("title", "Quarterly Business Review"),
        ("subtitle", "Nairobi office, Q3 2026"),
    ]


def test_bullet_levels_are_kept(rich):
    body = [t for t in rich.slides[1].texts if t.role == "body"]
    assert [(t.text, t.level) for t in body] == [
        ("Revenue grew 45% to KES 12.5m", 0),
        ("Driven by M-Pesa collections", 1),
        ("Churn: down to 3.2%", 0),
    ]


def test_body_placeholder_paragraphs_are_bulleted(rich):
    body = [t for t in rich.slides[1].texts if t.role == "body"]
    assert all(t.bullet == "char" for t in body)


def test_run_formatting_is_kept(rich):
    churn = rich.slides[1].texts[-1]
    assert [(r.text, r.bold, r.italic) for r in churn.runs] == [
        ("Churn: ", True, None),
        ("down to 3.2%", None, True),
    ]


def test_text_in_group_shapes_is_extracted(rich):
    assert texts(rich.slides[5]) == ["Grouped note A", "Grouped note B"]


def test_paragraphs_from_the_same_box_share_a_box_id(rich):
    body = [t for t in rich.slides[1].texts if t.role == "body"]
    assert len({t.box for t in body}) == 1
    assert rich.slides[1].texts[0].box != body[0].box


def test_empty_paragraphs_are_skipped(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tf = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text_frame
    tf.text = "First"
    tf.add_paragraph().text = "   "
    tf.add_paragraph().text = "Second"
    prs.save(tmp_path / "d.pptx")
    deck = extract_deck(tmp_path / "d.pptx")
    assert texts(deck.slides[0]) == ["First", "Second"]


def test_line_breaks_inside_a_paragraph_are_kept_as_one_item(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tf = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text_frame
    tf.text = "Line one\vline two"
    prs.save(tmp_path / "d.pptx")
    t = extract_deck(tmp_path / "d.pptx").slides[0].texts[0]
    assert t.text == "Line one\vline two"
    assert [r.br for r in t.runs] == [False, True, False]


def test_title_is_detected_without_placeholders(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    body = s.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(2)).text_frame
    body.text = "Body copy that explains the idea in more detail"
    body.paragraphs[0].runs[0].font.size = Pt(18)
    head = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1)).text_frame
    head.text = "Market opportunity"
    head.paragraphs[0].runs[0].font.size = Pt(40)
    prs.save(tmp_path / "d.pptx")
    slide = extract_deck(tmp_path / "d.pptx").slides[0]
    assert [(t.role, t.text) for t in slide.texts] == [
        ("title", "Market opportunity"),
        ("body", "Body copy that explains the idea in more detail"),
    ]


def test_items_follow_reading_order_top_to_bottom(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    # added bottom-first so z-order differs from reading order
    s.shapes.add_textbox(Inches(1), Inches(5), Inches(8), Inches(1)).text_frame.text = "Third"
    s.shapes.add_textbox(Inches(5), Inches(3), Inches(4), Inches(1)).text_frame.text = "Second right"
    s.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1)).text_frame.text = "Second left"
    prs.save(tmp_path / "d.pptx")
    slide = extract_deck(tmp_path / "d.pptx").slides[0]
    assert texts(slide) == ["Second left", "Second right", "Third"]


def test_footer_placeholder_text_is_role_footer(tmp_path):
    from pptx.enum.shapes import PP_PLACEHOLDER

    prs = Presentation()
    layout = prs.slide_layouts[1]
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = "Hello"
    footer_ph = next(p for p in layout.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.FOOTER)
    s.shapes.clone_placeholder(footer_ph)
    ft = next(p for p in s.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.FOOTER)
    ft.text = "Confidential - Acme Ltd"
    prs.save(tmp_path / "d.pptx")
    slide = extract_deck(tmp_path / "d.pptx").slides[0]
    assert ("footer", "Confidential - Acme Ltd") in [(t.role, t.text) for t in slide.texts]


# --- tables ---------------------------------------------------------------

def test_table_cells_and_positions(rich):
    (tbl,) = rich.slides[2].tables
    assert (tbl.rows, tbl.cols) == (4, 3)
    assert tbl.cell(1, 0).text == "Nairobi"
    assert tbl.cell(3, 2).text == "415"


def test_table_merges_are_recorded(rich):
    (tbl,) = rich.slides[2].tables
    origin = tbl.cell(0, 1)
    assert origin.text == "Units sold"
    assert (origin.row_span, origin.col_span) == (1, 2)
    assert tbl.cell(0, 2).spanned is True


# --- charts ---------------------------------------------------------------

def test_category_chart_data(rich):
    (chart,) = rich.slides[3].charts
    assert chart.chart_type == "COLUMN_CLUSTERED"
    assert chart.title == "Revenue (KES m)"
    assert chart.categories == ["Q1", "Q2", "Q3", "Q4"]
    assert [(s.name, s.values) for s in chart.series] == [
        ("2025", [10.5, 12.0, 14.25, 15.0]),
        ("2026", [13.0, 16.5, 21.0, None]),
    ]


def test_xy_chart_keeps_x_and_y_values(rich):
    (chart,) = rich.slides[5].charts
    assert chart.chart_type == "XY_SCATTER"
    (ser,) = chart.series
    assert ser.name == "Price vs volume"
    assert ser.x_values == [1.5, 2.0, 3.25]
    assert ser.values == [200.0, 180.0, 120.0]


# --- images ---------------------------------------------------------------

def test_image_binary_and_crop(rich):
    (img,) = rich.slides[4].images
    assert img.sha1 == hashlib.sha1(png_bytes(400, 300)).hexdigest()
    assert img.blob == png_bytes(400, 300)
    assert (img.px_width, img.px_height) == (400, 300)
    assert img.crop == pytest.approx((0.1, 0.0, 0.1, 0.0))


# --- notes and slide flags ------------------------------------------------

def test_notes_are_verbatim(rich):
    assert rich.slides[0].notes == "Welcome everyone.\nKeep intro under 2 minutes."
    assert rich.slides[1].notes == "Mention the Mombasa launch."
    assert rich.slides[2].notes == ""


def test_hidden_slide_is_flagged(rich):
    assert [s.hidden for s in rich.slides] == [False] * 6 + [True]


# --- SmartArt and unsupported content -------------------------------------

SMARTART_DATA = b"""<?xml version="1.0" encoding="UTF-8"?>
<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"
               xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <dgm:ptLst>
    <dgm:pt modelId="0" type="doc"><dgm:t><a:p><a:endParaRPr/></a:p></dgm:t></dgm:pt>
    <dgm:pt modelId="1"><dgm:t><a:p><a:r><a:t>Plan</a:t></a:r></a:p></dgm:t></dgm:pt>
    <dgm:pt modelId="2"><dgm:t><a:p><a:r><a:t>Budget KES 2m</a:t></a:r></a:p></dgm:t></dgm:pt>
    <dgm:pt modelId="3"><dgm:t><a:p><a:r><a:t>Execute</a:t></a:r></a:p></dgm:t></dgm:pt>
    <dgm:pt modelId="9" type="pres"><dgm:t><a:p><a:r><a:t>ignored</a:t></a:r></a:p></dgm:t></dgm:pt>
    <dgm:pt modelId="7" type="parTrans"/>
  </dgm:ptLst>
  <dgm:cxnLst>
    <dgm:cxn modelId="10" srcId="0" destId="1" srcOrd="0" destOrd="0"/>
    <dgm:cxn modelId="11" srcId="1" destId="2" srcOrd="0" destOrd="0"/>
    <dgm:cxn modelId="12" srcId="0" destId="3" srcOrd="1" destOrd="0"/>
    <dgm:cxn modelId="13" type="presOf" srcId="1" destId="9" srcOrd="0" destOrd="0"/>
  </dgm:cxnLst>
</dgm:dataModel>"""


def test_smartart_text_and_levels():
    assert parse_smartart_data(SMARTART_DATA) == [
        ("Plan", 0),
        ("Budget KES 2m", 1),
        ("Execute", 0),
    ]


def test_video_is_reported_as_unsupported(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    movie = tmp_path / "clip.mp4"
    movie.write_bytes(b"\x00\x00\x00\x18ftypmp42" + b"\x00" * 100)
    s.shapes.add_movie(str(movie), Inches(1), Inches(1), Inches(4), Inches(3), mime_type="video/mp4")
    prs.save(tmp_path / "d.pptx")
    deck = extract_deck(tmp_path / "d.pptx")
    assert deck.unsupported == ["Slide 1: video or audio"]
    assert deck.slides[0].images == []
