"""In-place rendering (Amendment 1B), locked objects (1C), and images (1E)."""

import io

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from deckbuilder import build_rich_deck, photo_bytes
from lockedbuilder import build_images_deck, build_locked_deck
from polisher.extract import extract_deck
from polisher.heuristic import plan_deck
from polisher.plan import parse_plan
from polisher.render.inplace import element_box, find_shape
from polisher.render.renderer import render
from polisher.render.themes import THEMES, get_theme
from polisher.verify import verify

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def polish(tmp_path, src, plan=None, theme="navy"):
    deck = extract_deck(src)
    out = tmp_path / f"out-{theme}.pptx"
    report = render(plan or plan_deck(deck), deck, get_theme(theme), out, src)
    return deck, out, report


def checked(deck, out, report):
    result = verify(deck, extract_deck(out), slide_sources=report.slide_sources,
                    dropped_backgrounds=report.dropped_backgrounds)
    assert result.passed, result.problems
    return result


@pytest.fixture(scope="module")
def locked_src(tmp_path_factory):
    return build_locked_deck(tmp_path_factory.mktemp("rl") / "locked.pptx")


@pytest.fixture(scope="module")
def images_src(tmp_path_factory):
    return build_images_deck(tmp_path_factory.mktemp("ri") / "images.pptx")


# --- locked objects ---------------------------------------------------------------

@pytest.mark.parametrize("theme", sorted(THEMES))
def test_locked_deck_round_trips_in_every_theme(tmp_path, locked_src, theme):
    checked(*polish(tmp_path, locked_src, theme=theme))


def test_video_animation_still_targets_the_video(tmp_path, locked_src):
    deck, out, _ = polish(tmp_path, locked_src)
    video_id = deck.slides[1].locked[0].shape_id
    slide_el = Presentation(out).slides[1]._element
    spids = slide_el.xpath(".//p:timing//p:spTgt/@spid")
    assert str(video_id) in spids


def test_locked_objects_keep_their_aspect_ratio(tmp_path, locked_src):
    deck, out, _ = polish(tmp_path, locked_src)
    prs = Presentation(out)
    for s, slide in zip(deck.slides[1:], list(prs.slides)[1:]):
        x = s.locked[0]
        el = find_shape(slide.shapes._spTree, x.shape_id)
        _, _, w, h = element_box(el)
        assert w / h == pytest.approx(x.aspect, rel=0.01), x.lock_type


def test_equations_are_never_enlarged(tmp_path, locked_src):
    before = element_box(find_shape(Presentation(locked_src).slides[5].shapes._spTree,
                                    extract_deck(locked_src).slides[5].locked[0].shape_id))
    deck, out, _ = polish(tmp_path, locked_src)
    after = element_box(find_shape(Presentation(out).slides[5].shapes._spTree, deck.slides[5].locked[0].shape_id))
    assert after[2] <= before[2] and after[3] <= before[3]


def test_media_parts_stay_in_the_package_once(tmp_path, locked_src):
    import zipfile

    _, out, _ = polish(tmp_path, locked_src)
    before, after = zipfile.ZipFile(locked_src).namelist(), zipfile.ZipFile(out).namelist()

    def parts(names, prefix):
        return sorted(n for n in names if n.startswith(prefix))

    for prefix in ("ppt/media/media", "ppt/embeddings/", "ppt/media/model3d", "ppt/ink/", "ppt/diagrams/"):
        assert parts(after, prefix) == parts(before, prefix), prefix


def test_animations_on_removed_text_are_pruned(tmp_path):
    src = build_rich_deck(tmp_path / "rich.pptx")
    prs = Presentation(src)
    slide = prs.slides[1]
    body_id = slide.placeholders[1].shape_id
    slide._element.append(etree.fromstring(
        f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        f'<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        f'<p:par><p:cTn id="3" fill="hold"><p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:childTnLst>'
        f'<p:par><p:cTn id="5" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect">'
        f'<p:childTnLst><p:set><p:cBhvr><p:cTn id="6" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{body_id}"/></p:tgtEl>'
        f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
        f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
        f'</p:childTnLst></p:cTn></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
        f'<p:bldLst><p:bldP spid="{body_id}" grpId="0" build="p"/></p:bldLst></p:timing>'
    ))
    prs.save(src)
    deck, out, report = polish(tmp_path, src)
    checked(deck, out, report)
    slide_el = Presentation(out).slides[1]._element
    assert not slide_el.xpath(f".//p:spTgt[@spid='{body_id}'] | .//p:bldP[@spid='{body_id}']")


# --- images ------------------------------------------------------------------------

@pytest.mark.parametrize("theme", sorted(THEMES))
def test_images_deck_round_trips_in_every_theme(tmp_path, images_src, theme):
    checked(*polish(tmp_path, images_src, theme=theme))


def test_dropped_background_is_logged(tmp_path, images_src):
    _, out, report = polish(tmp_path, images_src)
    assert report.dropped_backgrounds == ["s3-bg1"]
    assert any("background" in c for c in report.changes)
    assert not Presentation(out).slides[2]._element.xpath(".//p:bg//a:blip")


def test_background_can_be_kept_as_a_picture(tmp_path, images_src):
    deck = extract_deck(images_src)
    data = plan_deck(deck).to_dict()
    data["slides"][2] = {"source_slide": 3, "layout": "image_text",
                         "slots": {"title": ["s3-t1"], "visuals": ["s3-bg1"], "body": ["s3-t2"]}}
    deck, out, report = polish(tmp_path, images_src, plan=parse_plan(data))
    assert report.dropped_backgrounds == []
    checked(deck, out, report)
    out_deck = extract_deck(out)
    assert [i.sha256 for i in out_deck.slides[2].images] == [deck.slides[2].backgrounds[0].sha256]


def test_title_avoids_the_master_logo(tmp_path, images_src):
    _, out, _ = polish(tmp_path, images_src)
    slide = Presentation(out).slides[1]
    title = slide.shapes.title
    assert title.left + title.width <= Inches(8.6)


def test_split_slide_can_move_a_picture_across(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Cover"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Two photos"
    s.shapes.add_picture(io.BytesIO(photo_bytes(seed=0)), Inches(1), Inches(2), Inches(3))
    s.shapes.add_picture(io.BytesIO(photo_bytes(seed=1)), Inches(5), Inches(2), Inches(3))
    src = tmp_path / "two.pptx"
    prs.save(src)
    deck = extract_deck(src)
    plan = parse_plan({"slides": [
        {"source_slide": 1, "layout": "title", "slots": {"title": ["s1-t1"]}},
        {"source_slide": 2, "layout": "image_text", "slots": {"title": ["s2-t1"], "visuals": ["s2-i1"]}},
        {"source_slide": 2, "layout": "image_text", "slots": {"visuals": ["s2-i2"]}},
    ]})
    deck, out, report = polish(tmp_path, src, plan=plan)
    checked(deck, out, report)
    assert report.slide_sources == [1, 2, 2]
    assert len(Presentation(out).slides[2].shapes) >= 1


def test_wide_decks_stay_wide(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Wide cover"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Points"
    s.placeholders[1].text = "One point"
    src = tmp_path / "wide.pptx"
    prs.save(src)
    deck, out, report = polish(tmp_path, src)
    checked(deck, out, report)
    result = Presentation(out)
    assert result.slide_width == Inches(13.333)
    for slide in result.slides:
        for shape in slide.shapes:  # decoration may bleed off the edge; text never does
            if shape.has_text_frame and shape.text_frame.text.strip():
                assert shape.left + shape.width <= result.slide_width + Inches(0.01)


def test_decoration_never_covers_a_master_logo(tmp_path, images_src):
    _, out, _ = polish(tmp_path, images_src)
    logo = (Inches(8.6), Inches(0.2), Inches(1.2), Inches(0.6))
    for slide in Presentation(out).slides:
        for shape in slide.shapes:
            if shape.shape_type == 1 and not shape.has_text_frame or (shape.has_text_frame and not shape.text_frame.text):
                overlaps = (shape.left < logo[0] + logo[2] and logo[0] < shape.left + shape.width
                            and shape.top < logo[1] + logo[3] and logo[1] < shape.top + shape.height)
                assert not overlaps, shape.name
